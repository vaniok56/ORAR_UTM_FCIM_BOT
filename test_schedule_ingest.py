import sys
import os
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / "src"))

from schedule_ingest import Metadata, UploadReject, compare_versions, parse_title, parse_version, pdf_version, schedule_diff
from schedule_parser.audit_outputs import audit_schedule
from schedule_parser.parser import ScheduleBlock, SourceValue, _segments, parse_workbook, resolve_pair, resolve_pairs, write_schedule_workbook


class ScheduleIngestTests(unittest.TestCase):
    def dean_source(self, path, days=5, groups=("TI-261", "SI-222")):
        import openpyxl
        from openpyxl.styles import Border, Side
        book = openpyxl.Workbook()
        sheet = book.active
        sheet["A1"], sheet["B2"] = "ANUL UNIVERSITAR 2026/2027 ANUL II SEMESTRUL I", "Grupele"
        for column, group in enumerate(groups, 5):
            sheet.cell(2, column, group)
        for day_index, day in enumerate(("Luni", "Marţi", "Miercuri", "Joi", "Vineri", "Sâmbătă", "Duminică")[:days]):
            start = 3 + day_index * 43
            sheet.cell(start, 2, day)
            sheet.merge_cells(start_row=start, end_row=start + 41, start_column=2, end_column=2)
            for slot, time in enumerate(("8.00-9.30", "9.45-11.15", "11.30-13.00", "13.30-15.00", "15.15-16.45", "17.00-18.30", "18.45-20.15")):
                row = start + slot * 6
                sheet.cell(row, 3, time)
                sheet.merge_cells(start_row=row, end_row=row + 5, start_column=3, end_column=3)
                for column in range(5, 5 + len(groups)):
                    sheet.cell(row + 5, column).border = Border(bottom=Side(style="thin"))
        for offset, text in enumerate(("c. Algebra", "Example A.", "101", "c. Logic", "Sample B.", "102")):
            sheet.cell(3 + offset, 5, text)
        sheet["E6"].border = Border(top=Side(style="medium"))
        book.save(path)
        return book

    def test_upload_rejects_extent_merge_group_expansion_and_sunday(self):
        from unittest.mock import patch
        from schedule_ingest import MAX_SOURCE_ROWS, inspect_xlsx, prepare_upload
        from schedule_parser.parser import MAX_LOGICAL_GROUPS
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            for case, reason in (("row", "extent"), ("merge", "merged ranges"),
                                 ("groups", "Logical group count"), ("sunday", "Sunday")):
                with self.subTest(case=case):
                    book = self.dean_source(source, days=7 if case == "sunday" else 5)
                    if case == "row":
                        book.active.cell(MAX_SOURCE_ROWS + 1, 5, "too far")
                    elif case == "merge":
                        book.active.merge_cells("H1:AH1000")
                    elif case == "groups":
                        book.active["E2"] = "\n".join(f"TI-{100 + index}" for index in range(MAX_LOGICAL_GROUPS + 1))
                    book.save(source)
                    with self.assertRaisesRegex(ValueError, reason):
                        if case in {"row", "merge"}:
                            with patch("schedule_ingest.openpyxl.load_workbook", side_effect=AssertionError("must reject before load")):
                                inspect_xlsx(source)
                        else:
                            prepare_upload(source, None, 3, Path(directory) / "stage", None, None)

    def test_six_day_week_flows_through_parser_classifier_and_audit(self):
        import openpyxl
        from course_classification import build_classifications

        days6 = ("Luni", "Marţi", "Miercuri", "Joi", "Vineri", "Sâmbătă")
        times = ("8.00-9.30", "9.45-11.15", "11.30-13.00", "13.30-15.00",
                 "15.15-16.45", "17.00-18.30", "18.45-20.15")
        with tempfile.TemporaryDirectory() as directory:
            source, output = (Path(directory) / name for name in ("source.xlsx", "output.xlsx"))
            book = self.dean_source(source, days=6, groups=("TI-261",))
            sheet = book.active
            sheet["E254"] = "SO"  # Saturday final timeslot.
            sheet["E255"] = "Reițman P."
            sheet["E256"] = "101"
            book.save(source)

            blocks = parse_workbook(source)
            self.assertIn("Sâmbătă", {block.day for block in blocks})
            pairs = resolve_pairs(blocks)
            write_schedule_workbook(blocks, pairs, output)
            findings, _ = audit_schedule(source, output)
            self.assertEqual([item for item in findings if item.status != "approved"], [])

            written = openpyxl.load_workbook(output, data_only=True).active
            self.assertEqual(written.max_row, 1 + len(days6) * 7 * 2)
            self.assertEqual(written.cell(72, 1).value, "Sâmbătă")
            placed = {written.cell(row, column).value for row in range(72, written.max_row + 1)
                      for column in range(3, written.max_column + 1) if written.cell(row, column).value}
            self.assertIn("SO\nReițman P.\n101", placed)

            groups = [cell.value for cell in written[1][2:] if cell.value]
            labels, review = build_classifications(written, groups, 1, None, None, blocks, pairs)
            self.assertEqual(labels["SO\nReițman P.\n101"]["status"], "classified")
            self.assertFalse(any(row[2] == "Sâmbătă" for row in review))

    def test_classifier_reads_rows_beyond_the_old_seventy_one_row_cap(self):
        import openpyxl
        from course_classification import build_classifications

        book = openpyxl.Workbook()
        sheet = book.active
        sheet["C1"] = "TI-261"
        sheet["A85"] = "Sâmbătă"
        sheet["B85"] = "18.45-20.15"
        sheet["C85"] = "SO\nReițman P.\n101"
        labels, _ = build_classifications(sheet, ["TI-261"], 1, None, None)
        self.assertEqual(labels["SO\nReițman P.\n101"]["status"], "classified")

    def test_gap_values_flagged_but_empty_group_retained(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as directory:
            source, output = (Path(directory) / name for name in ("source.xlsx", "output.xlsx"))
            book = self.dean_source(source)
            sheet = book.active
            sheet.column_dimensions["F"].width = 1.285
            sheet["E45"] = "unexpected gap class"  # Between Monday and Tuesday.
            book.save(source)
            blocks = parse_workbook(source)
            write_schedule_workbook(blocks, resolve_pairs(blocks), output)
            findings, _ = audit_schedule(source, output)
            self.assertEqual({(item.source, item.reason) for item in findings}, {
                ("E45", "populated row outside six-row timeslots; inspect PDF ink"),
            })
            self.assertEqual(openpyxl.load_workbook(output).active["D1"].value, "SI-222")
            book = openpyxl.load_workbook(source)
            book.active["F3"] = "real group content"
            book.save(source)
            blocks = parse_workbook(source)
            write_schedule_workbook(blocks, resolve_pairs(blocks), output)
            findings, _ = audit_schedule(source, output)
            self.assertFalse(any(item.source == "F2" for item in findings))

    @unittest.skipUnless(os.environ.get("ORAR_PRIVATE_TESTS") == "1", "optional private geometry/source integration")
    def test_official_geometry_and_swapped_week_fixture_when_available(self):
        import csv
        import importlib.util

        root = Path.home() / "Downloads" / "Re__Orar (1)"
        if (not (root / "Anul_I_Semestrul_I.xlsx").exists() or not importlib.util.find_spec("pymupdf")
                or not (Path(__file__).parent / "schedule_parser" / "pdf_geometry.py").exists()):
            self.skipTest("official PDFs or optional development-only PyMuPDF unavailable")
        import openpyxl
        from schedule_parser.pdf_geometry import audit_geometry
        from course_classification import classification_counts
        from schedule_ingest import prepare_upload

        cases = (("Anul_I_Semestrul_I.xlsx", "anul_i_semestrul_i-27.pdf"),
                 ("Anul_II_Semestrul_III.xlsx", "anul_ii_semestrul_iii-18.pdf"),
                 ("Anul_III_Semestrul_V.xlsx", "anul_iii_semestrul_v-7.pdf"),
                 ("Anul_IV_Semestrul_VII.xlsx", "anul_iv_semestrul_vii-9.pdf"))
        expected_flags = (set(), {"L96", "V96", "W96", "Y96", "Z96", "AA96", "AB96"},
                          {"Z182"}, set())
        expected_labels = ((342, 1, 2), (219, 1, 1), (255, 0, 0), (210, 0, 0))
        totals = {}
        with tempfile.TemporaryDirectory() as directory:
            for year, (source, pdf) in enumerate(cases, 1):
                blocks = parse_workbook(root / source)
                output = Path(directory) / f"year{year}.xlsx"
                write_schedule_workbook(blocks, resolve_pairs(blocks), output)
                report = audit_geometry(root / source, root / pdf, output)
                stage = Path(directory) / f"stage{year}"
                arguments = (root / source, root / pdf, pdf_version(root / pdf), stage,
                             ["Luni", "Marţi", "Miercuri", "Joi", "Vineri"],
                             ["8.00-9.30", "9.45-11.15", "11.30-13.00", "13.30-15.00",
                              "15.15-16.45", "17.00-18.30", "18.45-20.15"])
                prepared = prepare_upload(*arguments)
                confirmed = sum(item["status"] == "classified" and not item.get("review_required")
                                for item in prepared.classifications.values())
                flagged = sum(bool(item.get("review_required")) for item in prepared.classifications.values())
                flagged_positions = {item[:5] for item in prepared.review if item[6] == "classified_review"}
                self.assertEqual((confirmed, flagged, len(flagged_positions)), expected_labels[year - 1])
                label_counts = classification_counts(prepared.classifications)
                self.assertEqual((label_counts["classified"], label_counts["needs_review"], label_counts["unclassified"]),
                                 (expected_labels[year - 1][0], expected_labels[year - 1][1], 0))
                self.assertFalse(any(item[6] != "classified_review" for item in prepared.review))
                self.assertEqual({item.source for item in prepared.findings}, expected_flags[year - 1])
                self.assertTrue(all(item.status == "approved" for item in prepared.findings))
                with (stage / "audit.csv").open(encoding="utf-8") as handle:
                    records = list(csv.DictReader(handle))
                    self.assertEqual({item["source"] for item in records}, expected_flags[year - 1])
                    self.assertTrue(all(item["status"] == "approved" for item in records))
                if expected_flags[year - 1]:
                    # Owner decision: the covered day-divider values are approved by
                    # coordinate and text, so they no longer depend on file hashes.
                    no_pdf, _ = audit_schedule(root / source, output)
                    self.assertTrue(all(item.status == "approved" for item in no_pdf))
                    other_pdf, _ = audit_schedule(root / source, output, root / cases[0][1])
                    self.assertTrue(all(item.status == "approved" for item in other_pdf
                                        if item.reason.startswith("covered day-divider")))
                if year == 2:
                    edited_source = Path(directory) / "edited_year2.xlsx"
                    book = openpyxl.load_workbook(root / source)
                    book.active["L96"] = "new gap text"
                    book.save(edited_source)
                    edited, _ = audit_schedule(edited_source, output, root / pdf)
                    self.assertEqual({item.source for item in edited if item.status == "unknown"}, {"L96"})
                for name, count in report["counts"].items():
                    totals[name] = totals.get(name, 0) + count
                self.assertEqual(report["counts"].get("unexplained_disagreement", 0), 0)
                self.assertEqual(report["counts"].get("merged_region_disagreement", 0), 0)
                self.assertFalse(any(item["type"] == "merged_half_disagreement" for item in report["findings"]))
                if year == 1:
                    book = openpyxl.load_workbook(output)
                    sheet = book.active
                    column = next(cell.column for cell in sheet[1] if cell.value == "IBM-261")
                    slot_index = 7 + 2  # Tuesday, third timeslot.
                    top = sheet.cell(2 + slot_index * 2, column)
                    bottom = sheet.cell(3 + slot_index * 2, column)
                    top.value, bottom.value = bottom.value, top.value
                    book.save(output)
                    swapped = audit_geometry(root / source, root / pdf, output)
                    self.assertTrue(any(item["group"] == "IBM-261" and item["day"] == "Marţi"
                                        and item["type"] == "unexplained_disagreement"
                                        for item in swapped["findings"]))
        self.assertEqual((totals["positions"], totals["exact_local_nonempty"], totals["exact_divided"],
                          totals["exact_local_empty"], totals["merged_region_unverified"], totals["merged_halves_contained"],
                          totals["continuous_two_slot_match"], totals["absent_pdf_header"],
                          totals["extracted_text_needs_raster"], totals["single_week_no_rule_needs_raster"]),
                         (4060, 834, 328, 2161, 952, 219, 34, 70, 8, 1))

    def test_top_only_divider_separates_week_parity(self):
        texts = ("AM", "Dohotaru L.", "720", "ALGA", None, "720")
        values = [SourceValue(67 + offset, 33, f"AG{67 + offset}", text, None,
                              "medium" if offset == 5 else None, "medium" if offset == 3 else None)
                  for offset, text in enumerate(texts)]
        segments = _segments(values)
        block = ScheduleBlock("IBM-261", "Marţi", "11.30-13.00", 67, 33, values, segments, [])
        pair = resolve_pair(block)
        self.assertEqual([(s.start_offset, s.end_offset) for s in segments], [(0, 2), (3, 5)])
        self.assertEqual((pair.odd_text, pair.even_text, pair.status),
                         ("AM\nDohotaru L.\n720", "ALGA\n720", "auto"))

    def test_source_border_audit_independently_rejects_combined_week_output(self):
        import openpyxl
        from openpyxl.styles import Border, Side

        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "dean.xlsx"
            output = Path(directory) / "output.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet["B1"] = "Grupele"
            sheet["E1"] = "IBM-261"
            times = ["8.00-9.30", "9.45-11.15", "11.30-13.00", "13.30-15.00",
                     "15.15-16.45", "17.00-18.30", "18.45-20.15"]
            for index, day in enumerate(("Luni", "Marţi", "Miercuri", "Joi", "Vineri")):
                for slot, time in enumerate(times):
                    row = 2 + (index * 7 + slot) * 6
                    sheet.cell(row, 2, day)
                    sheet.cell(row, 3, time)
            for offset, text in enumerate(("AM", "Dohotaru L.", "720", "ALGA", None, "720")):
                cell = sheet.cell(2 + offset, 5, text)
                if offset == 3:
                    cell.border = Border(top=Side(style="medium"))
                if offset == 5:
                    cell.border = Border(bottom=Side(style="medium"))
            sheet["E211"].border = Border(bottom=Side(style="thin"))
            workbook.save(source)
            blocks = parse_workbook(source)
            write_schedule_workbook(blocks, resolve_pairs(blocks), output)
            self.assertFalse(audit_schedule(source, output)[0])
            generated = openpyxl.load_workbook(output)
            generated.active["C2"] = "AM\nDohotaru L.\n720\nALGA\n720"
            generated.save(output)
            findings, _ = audit_schedule(source, output)
            self.assertTrue(any("top-only parity divider" in item.reason for item in findings))

    def test_sparse_midpoint_fragment_can_continue_only_into_next_upper_half(self):
        first_values = [SourceValue(100 + index, 5, f"E{100 + index}", text, None, None)
                        for index, text in enumerate((None, None, "lab. 0.5 gr.", None, "AFU", "Brînză M."))]
        second_values = [SourceValue(106 + index, 5, f"E{106 + index}", text, None,
                                     bottom_border="medium" if index == 2 else None)
                         for index, text in enumerate(("Țugulea V.", None, "419", None, None, None))]
        blocks = [
            ScheduleBlock("IBM-251", "Vineri", "9.45-11.15", 100, 5, first_values, _segments(first_values),
                          ["missing final border; checked for cross-timeslot continuation"]),
            ScheduleBlock("IBM-251", "Vineri", "11.30-13.00", 106, 5, second_values, _segments(second_values), []),
        ]
        pairs = resolve_pairs(blocks)
        self.assertTrue(all(pair.status == "auto" for pair in pairs))
        self.assertEqual(pairs[0].odd_text, "lab. 0.5 gr.\nAFU\nBrînză M.\nȚugulea V.\n419")
        self.assertEqual(pairs[0].odd_text, pairs[1].odd_text)

    def test_numbered_complete_upper_slot_does_not_autojoin_next_slot(self):
        def values(row, text):
            return [SourceValue(row + index, 5, f"E{row + index}", value, None,
                                bottom_border="medium" if index == 2 else None)
                    for index, value in enumerate(text)]
        first = values(100, ("1) lab. A", "Teacher A.", "101", None, None, None))
        second = values(106, ("2) lab. B", "Teacher B.", "102", None, None, None))
        blocks = [
            ScheduleBlock("TI-261", "Luni", "8.00-9.30", 100, 5, first, _segments(first), []),
            ScheduleBlock("TI-261", "Luni", "9.45-11.15", 106, 5, second, _segments(second), []),
        ]
        pairs = resolve_pairs(blocks)
        self.assertEqual(pairs[0].odd_text, "1) lab. A\nTeacher A.\n101")
        self.assertEqual(pairs[1].odd_text, "2) lab. B\nTeacher B.\n102")

    def test_top_border_inside_same_merge_does_not_split_weeks(self):
        values = [SourceValue(100 + index, 5, f"E{100 + index}", "CDE\nLitra D.\nA03",
                              "E100:E111", "medium" if index == 0 else None,
                              "medium" if index == 3 else None)
                  for index in range(6)]
        self.assertEqual([(segment.start_offset, segment.end_offset) for segment in _segments(values)], [(0, 5)])

    def test_title_requires_complete_dean_metadata(self):
        metadata = parse_title("ANUL UNIVERSITAR 2026/2027 ANUL II SEMESTRUL I")
        self.assertEqual(metadata, Metadata("2026/2027", 2, "I"))
        with self.assertRaises(UploadReject):
            parse_title("ANUL II SEMESTRUL I")

    def test_version_never_treats_unknown_as_final(self):
        self.assertEqual(parse_version("final"), "final")
        self.assertEqual(parse_version("12"), 12)
        with self.assertRaises(UploadReject):
            parse_version("0")
        with tempfile.TemporaryDirectory() as directory:
            self.assertIsNone(pdf_version(Path(directory) / "official.pdf"))
            self.assertEqual(pdf_version(Path(directory) / "orar_semestRul_I-12-copy.pdf"), 12)

    def test_version_order_requires_strictly_newer(self):
        self.assertEqual(compare_versions(3, 2), (True, "staged version is newer"))
        self.assertFalse(compare_versions(2, 2)[0])
        self.assertFalse(compare_versions(3, "final")[0])
        self.assertTrue(compare_versions("final", 99)[0])

    def test_diff_counts_parity_cells(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            paths = [Path(directory) / name for name in ("active.xlsx", "staged.xlsx")]
            for path, odd, even in ((paths[0], "A", "B"), (paths[1], "C", "B")):
                workbook = openpyxl.Workbook()
                sheet = workbook.active
                sheet.append([None, None, "TI-261"])
                sheet.append(["Luni", "8.00-9.30", odd])
                sheet.append([None, None, even])
                workbook.save(path)
            self.assertEqual(schedule_diff(*paths)[:3], (1, 0, 0))
            self.assertEqual(schedule_diff(*paths)[3][0][0], "changed")
            self.assertEqual(schedule_diff(*paths)[3][0][4], "ISO-even")
            revised = openpyxl.load_workbook(paths[1])
            revised.active["D1"] = "SI-261"
            revised.save(paths[1])
            self.assertTrue(any(row[0] == "group added" and row[3] == "SI-261"
                                for row in schedule_diff(*paths)[3]))


if __name__ == "__main__":
    unittest.main()
