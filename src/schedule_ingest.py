"""Validate and convert one dean schedule upload without publishing it."""

from __future__ import annotations

import csv
import hashlib
import re
import subprocess
import zipfile
from xml.parsers import expat
from openpyxl.utils import coordinate_to_tuple, range_boundaries
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl

from course_classification import build_classifications, review_csv, save_classifications
from schedule_groups import normalize_schedule_version
from schedule_parser.audit_outputs import audit_schedule, csv_cell, write_csv
from schedule_parser.parser import _displayed_value as displayed, parse_workbook, resolve_pairs, write_schedule_workbook


TITLE_RE = re.compile(
    r"ANUL UNIVERSITAR\s+(\d{4}/\d{4}).*?ANUL\s+([IVX]+).*?SEMESTRUL\s+([IVX]+)",
    re.IGNORECASE | re.DOTALL,
)
PDF_VERSION_RE = re.compile(r"semestrul[_ -]*[ivx]+-(\d{1,2})(?:-|$)", re.IGNORECASE)
ROMAN_YEAR = {"I": 1, "II": 2, "III": 3, "IV": 4}
PDF_MISSING_REASON = "source value not found by PDF text extraction"
MAX_XLSX_BYTES = 20 * 1024 * 1024
MAX_PDF_BYTES = 30 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
# Owner-approved from source XML measurements, 2026-10-03.
MAX_SOURCE_ROWS = 2000
MAX_SOURCE_COLUMNS = 1024
MAX_SOURCE_CELLS = 100_000
MAX_SOURCE_MERGES = 2000
MAX_SOURCE_MERGE_AREA = 20_000


class UploadReject(ValueError):
    """Upload cannot safely be published, including by override."""


@dataclass(frozen=True)
class Metadata:
    academic_year: str
    study_year: int
    semester: str


@dataclass(frozen=True)
class PreparedUpload:
    metadata: Metadata
    version: int | str
    xlsx: Path
    sidecar: Path
    audit_csv: Path
    review_csv: Path | None
    findings: tuple
    classifications: dict
    review: tuple


def parse_title(text: str) -> Metadata:
    match = TITLE_RE.search(" ".join(text.split()))
    if not match:
        raise UploadReject("title lacks academic year, study year, or semester")
    academic_year, study_year, semester = match.groups()
    if study_year.upper() not in ROMAN_YEAR:
        raise UploadReject("unsupported study year")
    return Metadata(academic_year, ROMAN_YEAR[study_year.upper()], semester.upper())


def _title_from_xlsx(path: Path) -> str:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if len(workbook.sheetnames) != 1:
        raise UploadReject("dean XLSX must contain exactly one worksheet")
    sheet = workbook.active
    header = next((row[0].row for row in sheet.iter_rows() if any(cell.value == "Grupele" for cell in row)), None)
    if header is None:
        raise UploadReject("dean XLSX lacks Grupele header")
    return " ".join(
        str(cell.value) for row in sheet.iter_rows(min_row=1, max_row=header - 1)
        for cell in row if cell.value is not None
    )


def inspect_xlsx(path: Path) -> Metadata:
    # parse also rejects fake/unsupported layout before an optional PDF is requested.
    validate_xlsx(path)
    parse_workbook(path)
    return parse_title(_title_from_xlsx(path))


def validate_xlsx(path: Path) -> None:
    if path.stat().st_size > MAX_XLSX_BYTES:
        raise UploadReject("XLSX exceeds 20 MiB")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            total = sum(item.file_size for item in entries)
            if len(entries) > 10_000 or total > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise UploadReject("XLSX archive exceeds safe limits")
            if "[Content_Types].xml" not in archive.namelist():
                raise UploadReject("file is not an XLSX workbook")
            if len({item.filename for item in entries}) != len(entries):
                raise UploadReject("duplicate XLSX archive part")
            worksheets = set()
            referenced = set()

            def reject_dtd(*args):
                raise UploadReject("XML DTD/entities are unsupported")

            def stream(name, start, end=None):
                parser = expat.ParserCreate(namespace_separator="}")
                parser.StartElementHandler = start
                parser.EndElementHandler = end
                parser.StartDoctypeDeclHandler = reject_dtd
                parser.EntityDeclHandler = reject_dtd
                parser.ExternalEntityRefHandler = reject_dtd
                with archive.open(name) as source:
                    while chunk := source.read(64 * 1024):
                        parser.Parse(chunk, False)
                    parser.Parse(b"", True)

            def content_type(tag, attrs):
                if tag.endswith("}Override") and attrs.get("ContentType", "").endswith("worksheet+xml"):
                    worksheets.add(attrs["PartName"].lstrip("/"))

            stream("[Content_Types].xml", content_type)
            sheet_count = 0

            def workbook_sheet(tag, attrs):
                nonlocal sheet_count
                if tag.endswith("}sheet"):
                    sheet_count += 1
                    if sheet_count > 1:
                        raise UploadReject("dean XLSX must contain exactly one worksheet")

            stream("xl/workbook.xml", workbook_sheet)
            if sheet_count != 1:
                raise UploadReject("dean XLSX must contain exactly one worksheet")
            # Check relationships too: worksheet parts need not be sheet1.xml or in xl/worksheets.
            import posixpath

            def relationship(tag, attrs):
                if tag.endswith("}Relationship") and attrs.get("Type", "").endswith("/worksheet"):
                    if attrs.get("TargetMode") == "External":
                        raise UploadReject("external worksheet is unsupported")
                    target = attrs["Target"]
                    referenced.add(posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target))

            stream("xl/_rels/workbook.xml.rels", relationship)
            if not referenced or not referenced <= worksheets or not worksheets <= set(archive.namelist()):
                raise UploadReject("missing or unsupported worksheet parts")
            if len(referenced) != 1:
                raise UploadReject("dean XLSX must contain exactly one worksheet")
            for name in worksheets:
                counts = {"cells": 0, "merges": 0, "area": 0}
                current_row = 0

                def extent(row, column):
                    if not 1 <= row <= MAX_SOURCE_ROWS or not 1 <= column <= MAX_SOURCE_COLUMNS:
                        raise UploadReject("worksheet row/column extent exceeds safe limits")

                def structure(tag, attrs):
                    nonlocal current_row
                    local = tag.rsplit("}", 1)[-1]
                    if local == "row":
                        if "r" not in attrs:
                            raise UploadReject("implicit row coordinates are unsupported")
                        current_row = int(attrs["r"])
                        extent(current_row, 1)
                    elif local == "c":
                        if not re.fullmatch(r"[A-Z]+[1-9]\d*", attrs.get("r", "")):
                            raise UploadReject("missing or invalid cell coordinates")
                        row, column = coordinate_to_tuple(attrs["r"])
                        extent(row, column)
                        if row != current_row:
                            raise UploadReject("cell coordinate differs from containing row")
                        counts["cells"] += 1
                        if counts["cells"] > MAX_SOURCE_CELLS:
                            raise UploadReject("worksheet cell count exceeds safe limits")
                    elif local == "mergeCell":
                        ref = attrs.get("ref", "")
                        if not re.fullmatch(r"[A-Z]+[1-9]\d*:[A-Z]+[1-9]\d*", ref):
                            raise UploadReject("unsupported merged range")
                        c1, r1, c2, r2 = range_boundaries(ref)
                        extent(r1, c1)
                        extent(r2, c2)
                        if r2 < r1 or c2 < c1:
                            raise UploadReject("invalid merged range")
                        counts["merges"] += 1
                        counts["area"] += (r2 - r1 + 1) * (c2 - c1 + 1)
                        if counts["merges"] > MAX_SOURCE_MERGES or counts["area"] > MAX_SOURCE_MERGE_AREA:
                            raise UploadReject("worksheet merged ranges exceed safe limits")

                def end_structure(tag):
                    nonlocal current_row
                    if tag.rsplit("}", 1)[-1] == "row":
                        current_row = 0

                stream(name, structure, end_structure)
            # Other XML parts (notably shared strings/styles) also must not expand entities.
            for name in archive.namelist():
                if name.endswith((".xml", ".rels")) and name not in worksheets:
                    stream(name, lambda tag, attrs: None)
    except zipfile.BadZipFile as error:
        raise UploadReject("file is not an XLSX workbook") from error
    except (expat.ExpatError, KeyError, TypeError, ValueError) as error:
        if isinstance(error, UploadReject):
            raise
        raise UploadReject("malformed or unsupported XLSX structure") from error


def inspect_pdf(path: Path) -> Metadata:
    if path.stat().st_size > MAX_PDF_BYTES:
        raise UploadReject("PDF exceeds 30 MiB")
    try:
        result = subprocess.run(
            ["pdftotext", "-enc", "UTF-8", "-raw", "-nopgbrk", str(path), "-"],
            check=True, capture_output=True, text=True, timeout=30,
        )
    except FileNotFoundError as error:
        raise UploadReject("PDF support unavailable: install pdftotext") from error
    except subprocess.CalledProcessError as error:
        raise UploadReject("cannot extract PDF text") from error
    except subprocess.TimeoutExpired as error:
        raise UploadReject("PDF text extraction exceeded 30 seconds") from error
    return parse_title(result.stdout)


def pdf_version(path: Path) -> int | None:
    match = PDF_VERSION_RE.search(path.stem)
    return int(match.group(1)) if match and 1 <= int(match.group(1)) <= 99 else None


def parse_version(value: str) -> int | str:
    value = value.strip().casefold()
    if value == "final":
        return "final"
    if value.isdigit() and 1 <= int(value) <= 99:
        return int(value)
    raise UploadReject("version must be 1-99 or final")


def prepare_upload(source: Path, pdf: Path | None, version: int | str, stage: Path, days, times) -> PreparedUpload:
    """Produce auditable staged files. Never touches active schedules."""
    metadata = inspect_xlsx(source)
    if pdf and metadata != inspect_pdf(pdf):
        raise UploadReject("PDF academic year, study year, or semester differs from dean XLSX")
    blocks = parse_workbook(source)
    if any(block.day == "Duminică" for block in blocks) or len({block.day for block in blocks}) > 6:
        raise UploadReject("Sunday/seven-day uploads are unsupported; runtime supports six days")
    pairs = resolve_pairs(blocks)
    if any(pair.status != "auto" for pair in pairs):
        raise UploadReject("parser could not determine every odd/even boundary")

    stage.mkdir(parents=True, exist_ok=True)
    output = stage / "schedule.xlsx"
    write_schedule_workbook(blocks, pairs, output, version=None if version == "final" else version, generated_date=date.today())
    try:
        findings, _ = audit_schedule(source, output, pdf)
    except ValueError as error:
        raise UploadReject(str(error)) from error
    audit_csv = stage / "audit.csv"
    write_csv(findings, audit_csv)
    hard_errors = [item for item in findings if item.status != "approved" and item.reason != PDF_MISSING_REASON]
    if hard_errors:
        raise UploadReject("source/output audit failed")

    workbook = openpyxl.load_workbook(output, data_only=True)
    schedule = workbook.active
    groups = [cell.value for cell in schedule[1][2:] if cell.value]
    classifications, review = build_classifications(schedule, groups, metadata.study_year, days, times, blocks, pairs)
    sidecar = stage / "schedule.classifications.json"
    save_classifications(output, classifications, sidecar, source)
    review_path = None
    if review:
        review_path = stage / "classification-review.csv"
        review_path.write_bytes(review_csv(review, hashlib.sha256(source.read_bytes()).hexdigest()))
    return PreparedUpload(metadata, version, output, sidecar, audit_csv, review_path, tuple(findings), classifications, tuple(review))


def compare_versions(staged: int | str, active_value) -> tuple[bool, str]:
    """Whether staged version is strictly newer than active. Advisory only.

    Owner decision 2026-09-29: publishing a same/older revision is allowed and
    this result is used for a warning, never to block publication.
    """
    try:
        active = normalize_schedule_version(active_value)
    except ValueError:
        return False, "active schedule version is invalid"
    if active == "final":
        return False, "active schedule is final"
    if staged == "final":
        return True, "staged final supersedes numeric active version"
    if staged > active:
        return True, "staged version is newer"
    return False, "staged version is same or older"


def schedule_diff(active_path: Path, staged_path: Path) -> tuple[int, int, int, list]:
    """Count and describe changed, added, removed parity cells without publishing."""

    def cells(path):
        sheet = openpyxl.load_workbook(path, data_only=True).active
        groups = [displayed(sheet, 1, column) for column in range(3, sheet.max_column + 1)]
        output, day = {}, ""
        for row in range(2, sheet.max_row, 2):
            day = (displayed(sheet, row, 1) or day).replace("ț", "ţ")
            time = displayed(sheet, row, 2)
            for column, group in enumerate(groups, 3):
                # Runtime uses isocalendar().week % 2; top row displays on ISO-even weeks.
                for parity, value_row in (("ISO-even", row), ("ISO-odd", row + 1)):
                    value = displayed(sheet, value_row, column)
                    if value:
                        output[day, time, group, parity] = value
        return output, set(filter(None, groups))

    (active, active_groups), (staged, staged_groups) = cells(active_path), cells(staged_path)
    differences = [
        ("added" if key not in active else "removed" if key not in staged else "changed", *key,
         active.get(key, ""), staged.get(key, ""))
        for key in sorted(active.keys() | staged.keys()) if active.get(key) != staged.get(key)
    ]
    differences.extend(("group added", "", "", group, "", "", "") for group in sorted(staged_groups - active_groups))
    differences.extend(("group removed", "", "", group, "", "", "") for group in sorted(active_groups - staged_groups))
    return (*(sum(row[0] == status for row in differences) for status in ("changed", "added", "removed")), differences)


def write_diff_csv(differences, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.writer(output)
        writer.writerow(("status", "day", "time", "group", "week", "active", "staged"))
        writer.writerows(tuple(csv_cell(value) for value in row) for row in differences)
