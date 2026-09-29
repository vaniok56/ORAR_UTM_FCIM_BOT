"""Manual terminal workflow for FCIM schedule parsing and validation."""

from __future__ import annotations

import csv
import hashlib
import json
import logging
import re
import shlex
import shutil
import subprocess
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Callable

import openpyxl

from audit_outputs import audit_schedule, write_csv, write_report
from parser import (
    parse_workbook,
    resolve_pairs,
    write_candidate_review_workbook,
    write_review_workbook,
    write_schedule_workbook,
)


TITLE_RE = re.compile(
    r"ANUL UNIVERSITAR\s+(\d{4}/\d{4}).*?ANUL\s+([IVX]+).*?SEMESTRUL\s+([IVX]+)",
    re.IGNORECASE | re.DOTALL,
)
PDF_VERSION_RE = re.compile(r"semestrul[_ -]*[ivx]+-(\d{1,2})(?:-|$)", re.IGNORECASE)


class Cancelled(Exception):
    pass


@dataclass(frozen=True)
class ScheduleMetadata:
    academic_year: str
    study_year: str
    semester: str
    groups: tuple[str, ...] = ()


@dataclass(frozen=True)
class ScheduleInput:
    xlsx: Path
    pdf: Path | None
    reference_xlsx: Path | None
    metadata: ScheduleMetadata
    version: int


@dataclass(frozen=True)
class ScheduleResult:
    label: str
    validation_scope: str
    version: int
    generated_date: str
    groups: int
    parity_cells: int
    warnings: int
    wrong: int
    missing: int
    extra: int
    unknown: int
    verdict: str
    output: Path


def parse_path(raw: str) -> Path:
    parts = shlex.split(raw.strip())
    if len(parts) != 1:
        raise ValueError("enter exactly one file path")
    return Path(parts[0]).expanduser().resolve()


def parse_paths(raw: str) -> list[Path]:
    parts = shlex.split(raw.strip())
    if not parts:
        raise ValueError("enter at least one file path")
    return [Path(part).expanduser().resolve() for part in parts]


def ask_file(
    prompt: str,
    suffix: str,
    optional: bool = False,
    input_fn: Callable[[str], str] = input,
    output_fn: Callable[[str], None] = print,
) -> Path | None:
    while True:
        raw = input_fn(f"{prompt}\n> ").strip()
        if raw.casefold() == "c":
            raise Cancelled
        if optional and not raw:
            return None
        if not raw:
            output_fn("Path is required.")
            continue
        try:
            path = parse_path(raw)
        except ValueError as error:
            output_fn(f"Invalid path: {error}")
            continue
        if path.is_dir():
            output_fn(f"Expected a {suffix} file, but this path is a folder.")
            continue
        if not path.is_file():
            output_fn("File does not exist.")
            continue
        if path.suffix.casefold() != suffix.casefold():
            output_fn(f"Expected a {suffix} file.")
            continue
        return path


def pdf_tools_available() -> bool:
    return bool(shutil.which("pdftotext") and shutil.which("pdftocairo"))


def parse_schedule_title(text: str) -> ScheduleMetadata:
    match = TITLE_RE.search(" ".join(text.split()))
    if not match:
        raise ValueError("schedule title does not contain academic year, study year, and semester")
    academic_year, study_year, semester = match.groups()
    return ScheduleMetadata(academic_year, study_year.upper(), semester.upper())


def read_workbook_title(path: Path) -> str:
    ws = openpyxl.load_workbook(path, data_only=True, read_only=False).active
    header_row = next(
        (row[0].row for row in ws.iter_rows() if any(cell.value == "Grupele" for cell in row)),
        None,
    )
    if header_row is None:
        raise ValueError("XLSX does not contain a 'Grupele' header")
    return " ".join(
        str(cell.value)
        for row in ws.iter_rows(min_row=1, max_row=header_row - 1)
        for cell in row
        if cell.value is not None
    )


def inspect_xlsx(path: Path) -> ScheduleMetadata:
    blocks = parse_workbook(path)
    metadata = parse_schedule_title(read_workbook_title(path))
    groups = tuple(dict.fromkeys(block.group for block in blocks))
    return ScheduleMetadata(metadata.academic_year, metadata.study_year, metadata.semester, groups)


def read_pdf_text(path: Path) -> str:
    result = subprocess.run(
        ["pdftotext", "-enc", "UTF-8", "-raw", "-nopgbrk", str(path), "-"],
        check=True,
        capture_output=True,
        text=True,
    )
    return result.stdout


def inspect_pdf(path: Path) -> ScheduleMetadata:
    return parse_schedule_title(read_pdf_text(path))


def parse_pdf_version(path: Path) -> int | None:
    match = PDF_VERSION_RE.search(path.stem)
    if not match:
        return None
    version = int(match.group(1))
    return version if 1 <= version <= 99 else None


def ask_version(pdf: Path | None, input_fn=input, output_fn=print) -> int:
    detected = parse_pdf_version(pdf) if pdf else None
    while True:
        choice = input_fn(
            f"Detected PDF version: {detected if detected is not None else 'unknown'}\n"
            + (f"Enter. Use version {detected}\n" if detected is not None else "Enter version explicitly.\n")
            + "Or type version [1-99] or final\n"
            "c. Cancel schedule\n> "
        ).strip().casefold()
        if not choice and detected is not None:
            return detected
        if choice == "c":
            raise Cancelled
        if choice == "final":
            return 0  # Legacy workbook encoding; never inferred from filename.
        if choice.isdigit() and 1 <= int(choice) <= 99:
            return int(choice)
        output_fn("Enter a number from 1 to 99, final, or c to cancel.")


def pair_mismatches(xlsx: ScheduleMetadata, pdf: ScheduleMetadata) -> list[str]:
    fields = ("academic_year", "study_year", "semester")
    return [field for field in fields if getattr(xlsx, field) != getattr(pdf, field)]


def revision_key(item: ScheduleInput):
    metadata = item.metadata
    return metadata.academic_year, metadata.study_year, metadata.semester


def schedule_label(metadata: ScheduleMetadata) -> str:
    academic_year = metadata.academic_year.replace("/", "_")
    return f"{academic_year}_year_{metadata.study_year}_semester_{metadata.semester}"


def sha256(path: Path) -> str:
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def create_run_directory(root: Path | None = None, now: datetime | None = None) -> Path:
    root = root or Path.home() / "Desktop" / "FCIM_schedule_runs"
    timestamp = (now or datetime.now()).strftime("%Y-%m-%d_%H-%M-%S-%f")
    path = root / timestamp
    path.mkdir(parents=True, exist_ok=False)
    return path


def copy_sources(item: ScheduleInput, destination: Path) -> tuple[ScheduleInput, dict]:
    sources = destination / "sources"
    sources.mkdir()
    copied_xlsx = Path(shutil.copy2(item.xlsx, sources / "dean.xlsx"))
    copied_pdf = Path(shutil.copy2(item.pdf, sources / "official.pdf")) if item.pdf else None
    copied_reference = (
        Path(shutil.copy2(item.reference_xlsx, sources / "transform_reference.xlsx"))
        if item.reference_xlsx
        else None
    )
    copies = ((item.xlsx, copied_xlsx), (item.pdf, copied_pdf), (item.reference_xlsx, copied_reference))
    for original, copied in copies:
        if original and copied and sha256(original) != sha256(copied):
            raise OSError(f"Copied source hash mismatch: {original}")
    manifest = {
        "label": schedule_label(item.metadata),
        "validation_scope": "xlsx+pdf" if item.pdf else "xlsx-only",
        "version": item.version,
        "original_xlsx": str(item.xlsx),
        "copied_xlsx": str(copied_xlsx.relative_to(destination.parent)),
        "xlsx_sha256": sha256(copied_xlsx),
        "original_pdf": str(item.pdf) if item.pdf else None,
        "copied_pdf": str(copied_pdf.relative_to(destination.parent)) if copied_pdf else None,
        "pdf_sha256": sha256(copied_pdf) if copied_pdf else None,
        "original_reference_xlsx": str(item.reference_xlsx) if item.reference_xlsx else None,
        "copied_reference_xlsx": (
            str(copied_reference.relative_to(destination.parent)) if copied_reference else None
        ),
        "reference_xlsx_sha256": sha256(copied_reference) if copied_reference else None,
    }
    copied_item = ScheduleInput(copied_xlsx, copied_pdf, copied_reference, item.metadata, item.version)
    return copied_item, manifest


def render_pdf(source: Path, output: Path):
    subprocess.run(
        ["pdftocairo", "-f", "1", "-l", "1", "-singlefile", "-png", "-r", "150", str(source), str(output.with_suffix(""))],
        check=True,
        capture_output=True,
    )


def process_schedule(
    item: ScheduleInput,
    output: Path,
    generated_date: date,
    progress_fn: Callable[[str], None] | None = None,
    manifest_fn: Callable[[dict], None] | None = None,
) -> tuple[ScheduleResult, dict]:
    if not 0 <= item.version <= 99:
        raise ValueError("Schedule version must be from 0 to 99")
    output.mkdir()
    copied, manifest = copy_sources(item, output)
    manifest["generated_date"] = generated_date.isoformat()
    if manifest_fn:
        manifest_fn(manifest)
    if progress_fn:
        progress_fn("[ok] Copied and verified sources")
    blocks = parse_workbook(copied.xlsx)
    if progress_fn:
        progress_fn(f"[ok] Loaded {len(blocks)} source blocks")
    pairs = resolve_pairs(blocks)
    write_review_workbook(blocks, output / "source_review.xlsx")
    write_schedule_workbook(
        blocks,
        pairs,
        output / "final_schedule.xlsx",
        version=copied.version,
        generated_date=generated_date,
    )
    if progress_fn:
        progress_fn("[ok] Wrote final schedule")
    write_candidate_review_workbook(blocks, copied.reference_xlsx, output / "candidate_review.xlsx")
    warnings = sum(bool(pair.review_flags) or pair.status != "auto" for pair in pairs)
    if progress_fn:
        progress_fn(f"[ok] Wrote {warnings} candidate warning rows")
    findings, stats = audit_schedule(copied.xlsx, output / "final_schedule.xlsx", copied.pdf)
    write_csv(findings, output / "audit.csv")
    write_report(findings, stats, output / "report.md")
    if copied.pdf:
        render_pdf(copied.pdf, output / "official.png")
    if progress_fn:
        progress_fn(f"[ok] Audit completed with {len(findings)} findings")
        if copied.pdf:
            progress_fn("[ok] Source/PDF text checks passed; geometry not checked" if all(item.status == "approved" for item in findings)
                        else "[!] Audit requires review")
        else:
            progress_fn("[-] Official PDF audit skipped")

    counts = Counter(item.status for item in findings)
    safe = all(item.status == "approved" for item in findings)
    scope = "xlsx+pdf" if copied.pdf else "xlsx-only"
    verdict = "SOURCE + PDF TEXT PASS (GEOMETRY NOT CHECKED)" if safe and copied.pdf else "SOURCE PASS (NO PDF)" if safe else "FAILED"
    result = ScheduleResult(
        label=schedule_label(item.metadata),
        validation_scope=scope,
        version=item.version,
        generated_date=generated_date.isoformat(),
        groups=stats["groups"],
        parity_cells=stats["output_parity_cells"],
        warnings=warnings,
        wrong=counts["wrong"],
        missing=counts["missing"],
        extra=counts["extra"],
        unknown=counts["unknown"],
        verdict=verdict,
        output=output,
    )
    return result, manifest


def write_session_files(
    run_directory: Path,
    results: list[ScheduleResult],
    manifests: list[dict],
    failures: list[str],
    created_at: datetime | None = None,
):
    payload = {
        "created_at": (created_at or datetime.now()).isoformat(),
        "schedules": manifests,
        "failures": failures,
    }
    (run_directory / "session.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    columns = [
        "label", "validation_scope", "version", "generated_date", "groups", "parity_cells", "warnings",
        "wrong", "missing", "extra", "unknown", "verdict",
    ]
    with (run_directory / "summary.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for result in results:
            row = asdict(result)
            writer.writerow({column: row[column] for column in columns})
    summary = ["# FCIM Schedule Lab run", ""]
    for result in results:
        summary.extend([
            f"## {result.label}",
            "",
            f"- Validation: {result.validation_scope}",
            f"- Version: {result.version}",
            f"- Generated date: {result.generated_date}",
            f"- Groups: {result.groups}",
            f"- Parity cells: {result.parity_cells}",
            f"- Warnings: {result.warnings}",
            f"- Wrong: {result.wrong}",
            f"- Missing: {result.missing}",
            f"- Extra: {result.extra}",
            f"- Unknown: {result.unknown}",
            f"- Verdict: {result.verdict}",
            "",
        ])
    if failures:
        summary.extend(["## Failures", "", *[f"- {failure}" for failure in failures], ""])
    (run_directory / "SUMMARY.md").write_text("\n".join(summary), encoding="utf-8")


def process_session(
    items: list[ScheduleInput],
    root: Path | None = None,
    progress_fn: Callable[[str], None] | None = None,
    started_at: datetime | None = None,
) -> tuple[Path, list[ScheduleResult], list[str]]:
    started_at = started_at or datetime.now()
    run_directory = create_run_directory(root, started_at)
    log_path = run_directory / "run.log"
    logger = logging.getLogger("schedule_parser.app")
    logger.handlers.clear()
    logger.setLevel(logging.INFO)
    handler = logging.FileHandler(log_path, encoding="utf-8")
    logger.addHandler(handler)
    results = []
    manifests = []
    failures = []
    for item in items:
        label = schedule_label(item.metadata)
        schedule_manifests = []
        if progress_fn:
            progress_fn(f"{label}: processing...")
        try:
            stage = (lambda message: progress_fn(f"{label}: {message}")) if progress_fn else None
            result, _ = process_schedule(
                item,
                run_directory / label,
                started_at.date(),
                stage,
                schedule_manifests.append,
            )
        except Exception:
            logger.exception("Processing failed for %s", label)
            failures.append(label)
            if progress_fn:
                progress_fn(f"{label}: failed. See run.log.")
        else:
            results.append(result)
            if progress_fn:
                progress_fn(f"{label}: {result.verdict}")
        manifests.extend(schedule_manifests)
    write_session_files(run_directory, results, manifests, failures, started_at)
    logger.removeHandler(handler)
    handler.close()
    return run_directory, results, failures


def confirm(prompt: str, input_fn=input) -> bool:
    return input_fn(f"{prompt} [Y/n]\n> ").strip().casefold() not in {"n", "no"}


def add_schedule_files(input_fn=input, output_fn=print) -> list[ScheduleInput]:
    raw = input_fn(
        "Drag or paste XLSX and optional matching PDF files for one or more study years.\n"
        "Files may be in any order. Do not enter a folder.\n> "
    ).strip()
    if raw.casefold() == "c":
        raise Cancelled
    paths = parse_paths(raw)
    if len(paths) != len(set(paths)):
        raise ValueError("same file path was entered more than once")
    for path in paths:
        if path.is_dir():
            raise ValueError(f"expected a file, but path is a folder: {path}")
        if not path.is_file():
            raise ValueError(f"file does not exist: {path}")
        if path.suffix.casefold() not in {".xlsx", ".pdf"}:
            raise ValueError(f"unsupported file type: {path}")

    xlsx_paths = [path for path in paths if path.suffix.casefold() == ".xlsx"]
    pdf_paths = [path for path in paths if path.suffix.casefold() == ".pdf"]
    if not xlsx_paths:
        raise ValueError("at least one XLSX file is required")
    if pdf_paths and not pdf_tools_available():
        raise ValueError(
            "PDF validation requires pdftotext and pdftocairo. "
            "Install Poppler with 'brew install poppler' or omit PDF files."
        )

    xlsx_by_key = {}
    for path in xlsx_paths:
        metadata = inspect_xlsx(path)
        key = metadata.academic_year, metadata.study_year, metadata.semester
        if key in xlsx_by_key:
            raise ValueError(f"multiple XLSX files identify {key}")
        xlsx_by_key[key] = path, metadata

    pdf_by_key = {}
    for path in pdf_paths:
        metadata = inspect_pdf(path)
        key = metadata.academic_year, metadata.study_year, metadata.semester
        if key in pdf_by_key:
            raise ValueError(f"multiple PDF files identify {key}")
        if key not in xlsx_by_key:
            raise ValueError(f"PDF has no matching XLSX: {path}")
        pdf_by_key[key] = path

    items = []
    for key, (xlsx, metadata) in xlsx_by_key.items():
        pdf = pdf_by_key.get(key)
        output_fn(
            f"Detected {metadata.academic_year}, Year {metadata.study_year}, "
            f"Semester {metadata.semester}: {len(metadata.groups)} groups, "
            f"PDF {'matched' if pdf else 'skipped'}"
        )
        items.append(ScheduleInput(xlsx, pdf, None, metadata, ask_version(pdf, input_fn, output_fn)))
    return items


def advanced_options(item: ScheduleInput, input_fn=input, output_fn=print) -> ScheduleInput:
    while True:
        reference = item.reference_xlsx.name if item.reference_xlsx else "none"
        choice = input_fn(
            f"Advanced options for Year {item.metadata.study_year}\n"
            f"Transform reference: {reference}\n\n"
            "1. Add or replace transform reference XLSX\n"
            "2. Remove transform reference\n"
            "Enter. Back\n> "
        ).strip()
        if not choice:
            return item
        if choice == "1":
            path = ask_file("Transform reference XLSX:", ".xlsx", input_fn=input_fn, output_fn=output_fn)
            item = ScheduleInput(item.xlsx, item.pdf, path, item.metadata, item.version)
        elif choice == "2":
            item = ScheduleInput(item.xlsx, item.pdf, None, item.metadata, item.version)
        else:
            output_fn("Choose 1, 2, or Enter.")


def review_session(items: list[ScheduleInput], output_fn=print):
    output_fn("\nSession review")
    for index, item in enumerate(items, 1):
        output_fn(
            f"\n{index}. {item.metadata.academic_year}, Year {item.metadata.study_year}, "
            f"Semester {item.metadata.semester}\n"
            f"   XLSX: {item.xlsx}\n"
            f"   PDF: {item.pdf or 'skipped'}\n"
            f"   Version: {item.version}\n"
            f"   Transform reference: {item.reference_xlsx or 'none'}"
        )


def open_path(path: Path):
    subprocess.run(["open", str(path)], check=True)


def run_new_session(
    input_fn=input,
    output_fn=print,
    root: Path | None = None,
    open_fn=open_path,
) -> tuple[int, bool]:
    items = []
    while True:
        output_fn("\nNew processing session")
        if items:
            review_session(items, output_fn)
        choice = input_fn(
            "\n1. Add schedule files\n"
            "2. Advanced options\n"
            "3. Process session\n"
            "4. Remove or replace schedule\n"
            "c. Cancel\n> "
        ).strip().casefold()
        if choice == "1":
            try:
                added_items = add_schedule_files(input_fn, output_fn)
            except (ValueError, subprocess.CalledProcessError) as error:
                output_fn(f"Could not inspect schedule: {error}")
                continue
            for item in added_items:
                existing = next(
                    (index for index, other in enumerate(items) if revision_key(other) == revision_key(item)),
                    None,
                )
                if existing is not None:
                    replace = input_fn(
                        "This academic year, study year, and semester is already selected.\n"
                        "1. Replace selected revision\n2. Keep selected revision\nc. Cancel\n> "
                    ).strip().casefold()
                    if replace == "1":
                        items[existing] = item
                    elif replace == "c":
                        raise Cancelled
                    elif replace != "2":
                        output_fn("Keeping selected revision.")
                else:
                    items.append(item)
        elif choice == "2":
            if not items:
                output_fn("Add a schedule first.")
                continue
            raw = input_fn(f"Choose schedule number [1-{len(items)}]:\n> ").strip()
            if raw.isdigit() and 1 <= int(raw) <= len(items):
                items[int(raw) - 1] = advanced_options(items[int(raw) - 1], input_fn, output_fn)
            else:
                output_fn("Invalid schedule number.")
        elif choice == "3":
            if not items:
                output_fn("Add at least one schedule first.")
                continue
            review_session(items, output_fn)
            if not confirm("Process selected schedules?", input_fn):
                continue
            run_directory, results, failures = process_session(items, root, output_fn)
            output_fn("\nProcessing complete")
            for result in results:
                output_fn(
                    f"{result.label}: {result.groups} groups, {result.warnings} warnings, "
                    f"{result.wrong} wrong, {result.missing} missing, {result.verdict}"
                )
            for failure in failures:
                output_fn(f"{failure}: FAILED. See run.log.")
            output_fn(f"Output: {run_directory}")
            session_code = 1 if failures or any(result.verdict == "FAILED" for result in results) else 0
            while True:
                action = input_fn(
                    "\n1. Open output folder\n2. Open summary\n"
                    "3. Open candidate review workbooks\nEnter. Return to menu\nq. Quit\n> "
                ).strip().casefold()
                if action in {"1", "2", "3"}:
                    if action == "3":
                        for result in results:
                            try:
                                open_fn(result.output / "candidate_review.xlsx")
                            except (OSError, subprocess.CalledProcessError) as error:
                                output_fn(f"Could not open file: {error}")
                    else:
                        try:
                            if action == "1":
                                open_fn(run_directory)
                            else:
                                open_fn(run_directory / "SUMMARY.md")
                        except (OSError, subprocess.CalledProcessError) as error:
                            output_fn(f"Could not open file: {error}")
                elif action == "q":
                    return session_code, True
                elif not action:
                    return session_code, False
                else:
                    output_fn("Choose 1, 2, 3, Enter, or q.")
        elif choice == "4":
            if not items:
                output_fn("No schedules selected.")
                continue
            raw = input_fn(f"Remove schedule number [1-{len(items)}]:\n> ").strip()
            if raw.isdigit() and 1 <= int(raw) <= len(items):
                items.pop(int(raw) - 1)
            else:
                output_fn("Invalid schedule number.")
        elif choice == "c":
            raise Cancelled
        else:
            output_fn("Choose 1, 2, 3, 4, or c.")


def validate_existing(input_fn=input, output_fn=print) -> int:
    source = ask_file("Source XLSX:", ".xlsx", input_fn=input_fn, output_fn=output_fn)
    generated = ask_file("Generated schedule XLSX:", ".xlsx", input_fn=input_fn, output_fn=output_fn)
    pdf = ask_file(
        "Drag or paste matching official PDF path. Press Enter without a path to skip PDF validation:",
        ".pdf",
        optional=True,
        input_fn=input_fn,
        output_fn=output_fn,
    )
    if pdf and not pdf_tools_available():
        output_fn("PDF validation requires pdftotext and pdftocairo. Install Poppler with 'brew install poppler'.")
        return 1
    if pdf:
        mismatches = pair_mismatches(inspect_xlsx(source), inspect_pdf(pdf))
        if mismatches:
            output_fn(f"PDF does not match source XLSX: {', '.join(mismatches)}")
            return 1
    csv_path = generated.with_name("audit.csv")
    report_path = generated.with_name("report.md")
    if (csv_path.exists() or report_path.exists()) and not confirm("Replace existing audit files?", input_fn):
        return 2
    findings, stats = audit_schedule(source, generated, pdf)
    write_csv(findings, csv_path)
    write_report(findings, stats, report_path)
    counts = Counter(item.status for item in findings)
    output_fn(
        f"Wrong: {counts['wrong']} Missing: {counts['missing']} "
        f"Extra: {counts['extra']} Unknown: {counts['unknown']}"
    )
    output_fn(f"Report: {report_path}")
    return 1 if any(item.status != "approved" for item in findings) else 0


def show_help(output_fn=print):
    output_fn(
        "Paste or drag each XLSX path manually. PDF is optional. "
        "No input directory is scanned. Each run copies selected sources "
        "into an immutable folder under Desktop/FCIM_schedule_runs."
    )


def run_wizard(input_fn=input, output_fn=print, root: Path | None = None, open_fn=open_path) -> int:
    exit_code = 0
    while True:
        choice = input_fn(
            "FCIM Schedule Lab\n\n"
            "1. New processing session\n"
            "2. Validate existing output\n"
            "3. Help\n"
            "q. Quit\n\n> "
        ).strip().casefold()
        try:
            if choice == "1":
                session_code, should_exit = run_new_session(input_fn, output_fn, root, open_fn)
                exit_code = max(exit_code, session_code)
                if should_exit:
                    return exit_code
            if choice == "2":
                exit_code = max(exit_code, validate_existing(input_fn, output_fn))
            if choice == "3":
                show_help(output_fn)
            elif choice == "q":
                return exit_code
            elif choice not in {"1", "2"}:
                output_fn("Choose 1, 2, 3, or q.")
        except Cancelled:
            output_fn("Cancelled.")
            exit_code = max(exit_code, 2)
        except Exception as error:
            output_fn(f"Operation failed: {error}")
            exit_code = max(exit_code, 1)


def main() -> int:
    try:
        return run_wizard()
    except KeyboardInterrupt:
        print("\nCancelled.")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
