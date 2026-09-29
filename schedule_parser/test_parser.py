import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from parser import (
    LayoutError,
    ScheduleBlock,
    SourceSegment,
    SourceValue,
    _layout,
    parse_workbook,
    resolve_pair,
    resolve_pairs,
    write_candidate_review_workbook,
    write_json,
    write_pairs_json,
    write_review_workbook,
    write_schedule_workbook,
)
from audit_outputs import audit_schedule


ROOT = Path(__file__).resolve().parents[1]
SCHEDULES = ROOT / "transform schedule" / "schedules"
SOURCES = sorted(
    path for path in SCHEDULES.glob("*.xlsx")
    if not path.name.endswith("_converted.xlsx") and not path.name.startswith("~$")
)
CURRENT_SOURCES = {
    "II": Path.home() / "Downloads" / "Anul_II_Semestrul_III-2.xlsx",
    "III": Path.home() / "Downloads" / "Anul_III_Semestrul_V.xlsx",
    "IV": Path.home() / "Downloads" / "Anul_IV_Semestrul_VII.xlsx",
}
CURRENT_SOURCES_AVAILABLE = all(path.is_file() for path in CURRENT_SOURCES.values())


class ParserTests(unittest.TestCase):
    def test_centered_three_line_entry_maps_to_both_after_continuation_pass(self):
        values = [
            SourceValue(
                row=100 + offset,
                column=5,
                coordinate=f"E{100 + offset}",
                value=value,
                merged_range=None,
                bottom_border=None,
            )
            for offset, value in enumerate((None, None, "L. Engleză 1", "Șișianu A.", "707", None))
        ]
        block = ScheduleBlock(
            group="CR-263",
            day="Vineri",
            time="11.30-13.00",
            source_row=100,
            source_column=5,
            values=values,
            segments=[SourceSegment(0, 5, values)],
            review_flags=["missing final border; checked for cross-timeslot continuation"],
        )
        pair = resolve_pairs([block])[0]
        self.assertEqual(pair.status, "auto")
        self.assertEqual(pair.odd_text, "L. Engleză 1\nȘișianu A.\n707")
        self.assertEqual(pair.odd_text, pair.even_text)
        self.assertIn("centered complete three-line entry treated as both weeks", pair.review_flags)

    def test_numbered_three_line_fragment_still_joins_following_timeslot(self):
        pairs = {
            (pair.group, pair.day, pair.time): pair
            for pair in resolve_pairs(parse_workbook(SCHEDULES / "Anul_II_Semestrul_IV (2).xlsx"))
        }
        pair = pairs[("CR-242", "Vineri", "11.30-13.00")]
        self.assertEqual(pair.status, "auto")
        self.assertIn("1) lab. 0.5 gr. RC", pair.odd_text)
        self.assertIn("2) lab. PCN", pair.odd_text)
        self.assertEqual(pair.odd_text, pair.even_text)
        self.assertIn("numbered cross-timeslot entries joined from source fragments", pair.review_flags)

    def test_multiline_header_maps_each_group_to_shared_source_column(self):
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws["B1"] = "Grupele"
        ws["E1"] = "R-263\nAI-264"
        ws["F1"] = "CR-261"
        ws["G1"] = "Grupele"
        self.assertEqual(
            _layout(ws),
            (1, 2, [(5, "R-263"), (5, "AI-264"), (6, "CR-261")]),
        )

    def test_known_workbooks_have_expected_block_counts(self):
        expected_groups = {
            "Anul_I_Semestrul_II (2).xlsx": 35,
            "Anul_II_Semestrul_IV (2).xlsx": 29,
            "Anul_III_2025_Semestrul_VI 2 (2).xlsx": 27,
            "Anul_IV_2025_Semestrul_VIII.xlsx": 4,
        }
        self.assertEqual({path.name for path in SOURCES}, set(expected_groups))
        for source in SOURCES:
            blocks = parse_workbook(source)
            self.assertEqual(len(blocks), expected_groups[source.name] * 35)
            self.assertEqual({block.time for block in blocks}, {
                "8.00-9.30", "9.45-11.15", "11.30-13.00", "13.30-15.00",
                "15.15-16.45", "17.00-18.30", "18.45-20.15",
            })

    def test_duplicate_weekly_timeslot_is_rejected(self):
        source = SCHEDULES / "Anul_II_Semestrul_IV (2).xlsx"
        workbook = openpyxl.load_workbook(source)
        ws = workbook.active
        header = next(cell for row in ws.iter_rows() for cell in row if cell.value == "Grupele")
        day_column = header.column
        time_column = header.column + 1
        timeslot_rows = [
            row for row in range(header.row + 1, ws.max_row + 1)
            if isinstance(ws.cell(row, time_column).value, str) and "-" in ws.cell(row, time_column).value
        ]
        day_merge = next(
            merged for merged in ws.merged_cells.ranges
            if merged.min_col <= day_column <= merged.max_col and merged.min_row <= timeslot_rows[1] <= merged.max_row
        )
        ws.unmerge_cells(str(day_merge))
        ws.cell(timeslot_rows[0], day_column).value = "Marți"
        ws.cell(timeslot_rows[1], day_column).value = "Marţi"
        ws.cell(timeslot_rows[1], time_column).value = ws.cell(timeslot_rows[0], time_column).value
        with tempfile.TemporaryDirectory() as directory:
            changed = Path(directory) / "duplicate.xlsx"
            workbook.save(changed)
            with self.assertRaisesRegex(LayoutError, "Duplicate weekly timeslot"):
                parse_workbook(changed)

    @unittest.skipUnless(
        (Path.home() / "Downloads" / "Re_ Orar" / "Anul_I_Semestrul_I.xlsx").is_file(),
        "current year I schedule is unavailable",
    )
    def test_current_year_one_shared_columns_are_complete(self):
        source = Path.home() / "Downloads" / "Re_ Orar" / "Anul_I_Semestrul_I.xlsx"
        blocks = parse_workbook(source)
        groups = list(dict.fromkeys(block.group for block in blocks))
        self.assertEqual(len(groups), 42)
        self.assertEqual(len(blocks), 42 * 35)
        self.assertEqual(groups[23:25], ["R-263", "AI-264"])
        self.assertEqual(groups[31:33], ["AI-263", "R-264"])
        self.assertIn("SI-222", groups)
        pairs = resolve_pairs(blocks)
        self.assertTrue(all(pair.status == "auto" for pair in pairs))
        cr263 = next(
            pair for pair in pairs
            if (pair.group, pair.day, pair.time) == ("CR-263", "Vineri", "11.30-13.00")
        )
        self.assertEqual(cr263.odd_text, "L. Engleză 1\n707")
        self.assertEqual(cr263.odd_text, cr263.even_text)
        self.assertTrue(cr263.review_flags)

    def test_same_sized_unrelated_reference_is_rejected(self):
        blocks = parse_workbook(SCHEDULES / "Anul_II_Semestrul_IV (2).xlsx")
        groups = list(dict.fromkeys(block.group for block in blocks))
        times = list(dict.fromkeys((block.day, block.time, block.source_row) for block in blocks))
        reference = openpyxl.Workbook()
        reference.active.cell(len(times) * 2 + 1, len(groups) + 2).value = "wrong schedule"
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            reference_path = root / "reference.xlsx"
            reference.save(reference_path)
            with self.assertRaisesRegex(LayoutError, "groups do not match"):
                write_candidate_review_workbook(blocks, reference_path, root / "candidate.xlsx")

    def test_multi_instructor_shape_maps_to_both_weeks(self):
        source = SCHEDULES / "Anul_II_Semestrul_IV (2).xlsx"
        block = next(
            block for block in parse_workbook(source)
            if (block.group, block.day, block.time) == ("IBM-241", "Miercuri", "9.45-11.15")
        )
        self.assertEqual(
            [value.value for value in block.values],
            [None, "c. Biomateriale", "Pocaznoi I.", "Duda B.", "405", None],
        )
        self.assertEqual([(segment.start_offset, segment.end_offset) for segment in block.segments], [(0, 5)])
        pair = resolve_pair(block)
        self.assertEqual(pair.status, "auto")
        self.assertEqual(
            pair.odd_text,
            "c. Biomateriale\nPocaznoi I.\nDuda B.\n405",
        )
        self.assertEqual(pair.odd_text, pair.even_text)

    def test_two_border_separated_entries_map_to_odd_and_even(self):
        source = SCHEDULES / "Anul_II_Semestrul_IV (2).xlsx"
        block = next(
            block for block in parse_workbook(source)
            if (block.group, block.day, block.time) == ("IBM-241", "Luni", "13.30-15.00")
        )
        self.assertEqual(
            [value.value for value in block.values],
            ["lab. CI", "Magariu N.", "406", "lab. ME", "Lupan C.", "422"],
        )
        self.assertEqual([(segment.start_offset, segment.end_offset) for segment in block.segments], [(0, 2), (3, 5)])
        pair = resolve_pair(block)
        self.assertEqual(pair.status, "auto")
        self.assertEqual(pair.odd_text, "lab. CI\nMagariu N.\n406")
        self.assertEqual(pair.even_text, "lab. ME\nLupan C.\n422")

    def test_half_pair_entries_map_to_their_border_indicated_week(self):
        source = SCHEDULES / "Anul_III_2025_Semestrul_VI 2 (2).xlsx"
        odd = next(
            block for block in parse_workbook(source)
            if (block.group, block.day, block.time) == ("SI-231", "Luni", "17.00-18.30")
        )
        even = next(
            block for block in parse_workbook(source)
            if (block.group, block.day, block.time) == ("IBM-231", "Luni", "9.45-11.15")
        )
        self.assertEqual(resolve_pair(odd).odd_text, "c. PMRI\nPutere A.\n505")
        self.assertIsNone(resolve_pair(odd).even_text)
        self.assertIsNone(resolve_pair(even).odd_text)
        self.assertEqual(resolve_pair(even).even_text, "lab. DMDT1\nȚugulea V.\n422")

    def test_unbordered_six_rows_split_into_two_complete_entries(self):
        source = SCHEDULES / "Anul_I_Semestrul_II (2).xlsx"
        block = next(
            block for block in parse_workbook(source)
            if (block.group, block.day, block.time) == ("TI-251", "Luni", "8.00-9.30")
        )
        pair = resolve_pair(block)
        self.assertEqual(pair.status, "auto")
        self.assertEqual(pair.odd_text, "AM\nCostaș A.\n202")
        self.assertEqual(pair.even_text, "c. Analiza Matematică 2\nCostaș A.\n3-3")

    def test_duplicate_line_layout_maps_to_one_entry_for_both_weeks(self):
        source = SCHEDULES / "Anul_IV_2025_Semestrul_VIII.xlsx"
        block = next(
            block for block in parse_workbook(source)
            if (block.group, block.day, block.time) == ("FAF-221", "Luni", "15.15-16.45")
        )
        pair = resolve_pair(block)
        self.assertEqual(pair.status, "auto")
        self.assertEqual(pair.odd_text, "c.  Testarea Software\nCatruc M.\n104")
        self.assertEqual(pair.odd_text, pair.even_text)

    def test_cross_timeslot_fragments_are_joined(self):
        source = SCHEDULES / "Anul_III_2025_Semestrul_VI 2 (2).xlsx"
        pairs = {
            (pair.group, pair.day, pair.time): pair
            for pair in resolve_pairs(parse_workbook(source))
        }
        expected = "lab. 0.5 gr.\nSM\nPalamarciuc N.\n524"
        for time in ("13.30-15.00", "15.15-16.45"):
            self.assertEqual(pairs[("AI-231", "Luni", time)].odd_text, expected)
            self.assertEqual(pairs[("AI-231", "Luni", time)].even_text, expected)

    def test_shifted_three_row_entry_maps_to_both_weeks(self):
        source = SCHEDULES / "Anul_III_2025_Semestrul_VI 2 (2).xlsx"
        block = next(
            block for block in parse_workbook(source)
            if (block.group, block.day, block.time) == ("FAF-231", "Marţi", "17.00-18.30")
        )
        pair = resolve_pair(block)
        self.assertEqual(pair.status, "auto")
        self.assertEqual(pair.odd_text, "c. Circuite și Dispozitive Electronice\nMagariu N.\n104")
        self.assertEqual(pair.odd_text, pair.even_text)

    def test_unbordered_complete_rows_keep_their_week_boundary(self):
        source = SCHEDULES / "Anul_III_2025_Semestrul_VI 2 (2).xlsx"
        block = next(
            block for block in parse_workbook(source)
            if (block.group, block.day, block.time) == ("TI-235", "Luni", "17.00-18.30")
        )
        pair = resolve_pair(block)
        self.assertEqual(pair.odd_text, "BD\nBulai R.\n606")
        self.assertEqual(pair.even_text, "TPP\nMititelu A.\n630")

    def test_numbered_cross_timeslot_entries_are_joined(self):
        source = SCHEDULES / "Anul_II_Semestrul_IV (2).xlsx"
        pairs = {
            (pair.group, pair.day, pair.time): pair
            for pair in resolve_pairs(parse_workbook(source))
        }
        expected = (
            "1) lab. 0.5 gr. RC\nSeniușin A.\n215\n\n"
            "2) lab. PCN\n0.5 gr.\nMunteanu S.\n407"
        )
        for time in ("11.30-13.00", "13.30-15.00"):
            self.assertEqual(pairs[("CR-242", "Vineri", time)].odd_text, expected)
            self.assertEqual(pairs[("CR-242", "Vineri", time)].even_text, expected)

    def test_ambiguous_cross_timeslot_tail_requires_review(self):
        source = SCHEDULES / "Anul_II_Semestrul_IV (2).xlsx"
        pairs = {
            (pair.group, pair.day, pair.time): pair
            for pair in resolve_pairs(parse_workbook(source))
        }
        # This older workbook has no PDF proof for whether the trailing room
        # belongs across the slot boundary. Do not manufacture a complete class.
        self.assertEqual(pairs[("SI-241", "Miercuri", "17.00-18.30")].status, "review")
        self.assertEqual(pairs[("SI-241", "Miercuri", "18.45-20.15")].odd_text, "114")

    def test_exports_are_reopenable(self):
        blocks = parse_workbook(SOURCES[0])
        with tempfile.TemporaryDirectory() as directory:
            json_path = Path(directory) / "blocks.json"
            pairs_json_path = Path(directory) / "pairs.json"
            xlsx_path = Path(directory) / "review.xlsx"
            schedule_xlsx_path = Path(directory) / "schedule.xlsx"
            candidates_xlsx_path = Path(directory) / "candidates.xlsx"
            write_json(blocks, json_path)
            write_pairs_json(resolve_pairs(blocks), pairs_json_path)
            write_review_workbook(blocks, xlsx_path)
            write_schedule_workbook(blocks, resolve_pairs(blocks), schedule_xlsx_path)
            write_candidate_review_workbook(
                blocks,
                SCHEDULES / "Anul_III_2025_Semestrul_VI 2 (2)_converted.xlsx",
                candidates_xlsx_path,
            )
            self.assertTrue(json_path.read_text(encoding="utf-8").startswith("["))
            self.assertTrue(pairs_json_path.read_text(encoding="utf-8").startswith("["))
            self.assertEqual(openpyxl.load_workbook(xlsx_path).active.max_row, len(blocks) + 1)
            self.assertEqual(openpyxl.load_workbook(schedule_xlsx_path).active.max_row, 71)
            schedule_sheet = openpyxl.load_workbook(schedule_xlsx_path).active
            self.assertEqual(schedule_sheet.cell(16, 3).border.top.style, "medium")
            candidate_sheet = openpyxl.load_workbook(candidates_xlsx_path).active
            self.assertEqual(
                candidate_sheet.max_row,
                sum(pair.status == "review" or bool(pair.review_flags) for pair in resolve_pairs(blocks)) + 1,
            )
            self.assertEqual(candidate_sheet.cell(1, 11).value, "resolution status")
            self.assertEqual(candidate_sheet.cell(1, 14).value, "old odd candidate")

    @unittest.skipUnless(CURRENT_SOURCES_AVAILABLE, "current schedule workbooks are not available")
    def test_current_two_line_entry_applies_to_both_weeks(self):
        pair = next(
            pair for pair in resolve_pairs(parse_workbook(CURRENT_SOURCES["II"]))
            if (pair.group, pair.day, pair.time) == ("TI-255", "Marţi", "15.15-16.45")
        )
        self.assertEqual(pair.status, "auto")
        self.assertEqual(pair.odd_text, "L. Engleză\n312")
        self.assertEqual(pair.even_text, "L. Engleză\n312")

    @unittest.skipUnless(CURRENT_SOURCES_AVAILABLE, "current schedule workbooks are not available")
    def test_current_mixed_parity_entries_follow_source_border(self):
        pair = next(
            pair for pair in resolve_pairs(parse_workbook(CURRENT_SOURCES["II"]))
            if (pair.group, pair.day, pair.time) == ("CR-251", "Luni", "9.45-11.15")
        )
        self.assertEqual(pair.odd_text, "lab. CDE\nChiriac M.\nA03")
        self.assertEqual(pair.even_text, "L. Engleză\n606")

    @unittest.skipUnless(CURRENT_SOURCES_AVAILABLE, "current schedule workbooks are not available")
    def test_current_numbered_cross_timeslot_fragments_are_joined(self):
        pairs = {
            (pair.group, pair.day, pair.time): pair
            for pair in resolve_pairs(parse_workbook(CURRENT_SOURCES["II"]))
        }
        expected = "lab. 0.5 gr.\n1) CDE\nChiriac M.\n\nA03\n2) MS\nLitra D.\n422"
        for time in ("8.00-9.30", "9.45-11.15"):
            self.assertEqual(pairs[("IBM-251", "Marţi", time)].odd_text, expected)
            self.assertEqual(pairs[("IBM-251", "Marţi", time)].even_text, expected)

    @unittest.skipUnless(CURRENT_SOURCES_AVAILABLE, "current schedule workbooks are not available")
    def test_current_four_line_entry_preserves_room(self):
        pair = next(
            pair for pair in resolve_pairs(parse_workbook(CURRENT_SOURCES["III"]))
            if (pair.group, pair.day, pair.time) == ("FAF-241", "Miercuri", "18.45-20.15")
        )
        self.assertEqual(pair.odd_text, "DAS\nPoștaru A.\nZaica M.\n113")
        self.assertEqual(pair.even_text, pair.odd_text)

    @unittest.skipUnless(CURRENT_SOURCES_AVAILABLE, "current schedule workbooks are not available")
    def test_current_multiline_merge_is_not_duplicated(self):
        pair = next(
            pair for pair in resolve_pairs(parse_workbook(CURRENT_SOURCES["IV"]))
            if (pair.group, pair.day, pair.time) == ("MN-232", "Miercuri", "11.30-13.00")
        )
        self.assertEqual(pair.odd_text, "lab. 0.5 gr TVLSIN\nTrofim V.\n426")
        self.assertEqual(pair.even_text, pair.odd_text)

    @unittest.skipUnless(CURRENT_SOURCES_AVAILABLE, "current schedule workbooks are not available")
    def test_current_schedules_have_no_review_placeholders(self):
        with tempfile.TemporaryDirectory() as directory:
            for year, source in CURRENT_SOURCES.items():
                blocks = parse_workbook(source)
                output = Path(directory) / f"year-{year}.xlsx"
                write_schedule_workbook(blocks, resolve_pairs(blocks), output)
                values = {
                    cell.value
                    for row in openpyxl.load_workbook(output, data_only=True).active.iter_rows()
                    for cell in row
                }
                self.assertNotIn("[REVIEW REQUIRED]", values)

    @unittest.skipUnless(CURRENT_SOURCES_AVAILABLE, "current schedule workbooks are not available")
    def test_older_current_sources_gaps_are_approved_by_coordinate(self):
        with tempfile.TemporaryDirectory() as directory:
            for year, source in CURRENT_SOURCES.items():
                blocks = parse_workbook(source)
                output = Path(directory) / f"year-{year}.xlsx"
                write_schedule_workbook(blocks, resolve_pairs(blocks), output)
                findings, _ = audit_schedule(source, output)
                self.assertTrue(findings, year) if year in {"II", "III"} else self.assertEqual(findings, [], year)
                self.assertTrue(all(item.status == "approved" for item in findings), year)

    def test_every_extracted_value_has_exact_source_provenance(self):
        for source in SOURCES:
            workbook = openpyxl.load_workbook(source, data_only=True)
            ws = workbook.active
            for block in parse_workbook(source):
                for value in block.values:
                    if value.merged_range:
                        merged = next(item for item in ws.merged_cells.ranges if str(item) == value.merged_range)
                        source_value = ws.cell(merged.min_row, merged.min_col).value
                    else:
                        source_value = ws.cell(value.row, value.column).value
                    expected = str(source_value).strip() if source_value is not None else None
                    self.assertEqual(value.value, expected, value.coordinate)

    def test_parsing_never_mutates_source_files(self):
        for source in SOURCES:
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            parse_workbook(source)
            after = hashlib.sha256(source.read_bytes()).hexdigest()
            self.assertEqual(after, before, source.name)

    def test_unknown_layout_fails_loudly(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "invalid.xlsx"
            openpyxl.Workbook().save(source)
            with self.assertRaises(LayoutError):
                parse_workbook(source)


if __name__ == "__main__":
    unittest.main()
