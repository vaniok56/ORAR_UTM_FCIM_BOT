"""Development-only PDF grid comparison. Never a publication/visibility oracle.

Install PyMuPDF separately to run: python -m schedule_parser.pdf_geometry
  --source dean.xlsx --pdf dean.pdf --output staged_schedule.xlsx
Covered text remains extractable; every disagreement requires raster inspection.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from .audit_outputs import displayed_value
from .parser import _layout, _timeslots


def tokens(value):
    return Counter(re.findall(r"[^\W_]+", unicodedata.normalize("NFC", str(value or ""))
                              .casefold().replace("ţ", "ț"), re.UNICODE))


def compare_half(top, bottom, pdf_top, pdf_bottom, divided):
    """Exact token comparison only when PDF shows an actual local divider."""
    upper, lower = tokens(top), tokens(bottom)
    if divided:
        return upper == tokens(pdf_top) and lower == tokens(pdf_bottom)
    return upper == lower and upper == tokens(pdf_top) + tokens(pdf_bottom)


def pdf_grid(page, columns, slots):
    words = page.get_text("words")
    time_re = re.compile(r"\d{1,2}\.\d{2}-\d{1,2}\.\d{2}")
    time_words = sorted((w for w in words if w[0] < 120 and time_re.fullmatch(w[4])), key=lambda w: w[1])
    if [w[4] for w in time_words] != [time for _, _, time in slots]:
        raise ValueError("PDF time grid differs from XLSX; manual review required")
    header_words = defaultdict(list)
    for word in words:
        if word[1] < time_words[0][1] - 3:
            header_words[word[4]].append(word)
    centers = defaultdict(list)
    for col, group in columns:
        hits = header_words[group]
        if len(hits) == 1:
            centers[col].append((hits[0][0] + hits[0][2]) / 2)
    physical = sorted((col, sum(xs) / len(xs)) for col, xs in centers.items())
    if len(physical) < 2 or any(a[1] >= b[1] for a, b in zip(physical, physical[1:])):
        raise ValueError("PDF group grid cannot be mapped unambiguously")
    bounds = {}
    for index, (col, center) in enumerate(physical):
        left = (physical[index - 1][1] + center) / 2 if index else center - (physical[1][1] - center) / 2
        right = ((physical[index + 1][1] + center) / 2 if index + 1 < len(physical)
                 else center + (center - physical[index - 1][1]) / 2)
        bounds[col] = (left, right)
    rules = [item[1] for drawing in page.get_drawings() for item in drawing["items"]
             if item[0] == "re" and item[1].height < 2 and item[1].width > 5]
    tops = []
    for word in time_words:
        center = (word[0] + word[2]) / 2
        candidates = [r.y0 for r in rules if r.x0 <= center - 1 and r.x1 >= center + 1
                      and 4 < word[1] - r.y0 < 13]
        if not candidates:
            raise ValueError(f"Missing PDF time rule at {word[4]}")
        tops.append(max(candidates))
    last = time_words[-1]
    center = (last[0] + last[2]) / 2
    finishes = [r.y0 for r in rules if r.x0 <= center - 1 and r.x1 >= center + 1
                and 4 < r.y0 - last[1] < 19]
    if not finishes:
        raise ValueError("Missing final PDF time rule")
    bottoms = tops[1:] + [min(finishes)]
    return words, bounds, tops, bottoms, rules


def audit_geometry(source_path, pdf_path, output_path):
    try:
        import pymupdf
    except ImportError as error:
        raise RuntimeError("Install PyMuPDF outside bot runtime for development geometry audit") from error

    ws = openpyxl.load_workbook(source_path, data_only=True).active
    header_row, day_col, groups = _layout(ws)
    slots = _timeslots(ws, header_row, day_col, day_col + 1)
    output = openpyxl.load_workbook(output_path, data_only=True).active
    if output.max_row != 71 or [displayed_value(output, 1, col) for col in range(3, output.max_column + 1)] != [name for _, name in groups]:
        raise ValueError("Output group inventory/grid differs from XLSX")
    with pymupdf.open(pdf_path) as document:
        if len(document) != 1:
            raise ValueError("PDF page count differs from supported single-page layout")
        words, bounds, tops, bottoms, rules = pdf_grid(document[0], groups, slots)

    counts = Counter()
    findings = []
    source_merges = list(ws.merged_cells.ranges)
    for group_index, (col, group) in enumerate(groups):
        if col not in bounds:
            counts["absent_pdf_header"] += len(slots)
            populated = [(day, time) for index, (_, day, time) in enumerate(slots)
                         if displayed_value(output, 2 + index * 2, 3 + group_index)
                         or displayed_value(output, 3 + index * 2, 3 + group_index)]
            findings.append({"type": "absent_nonempty_header" if populated else "absent_pdf_header",
                             "group": group, "source": ws.cell(header_row, col).coordinate,
                             "populated_slots": populated})
            continue
        left, right = bounds[col]
        for index, (row, day, time) in enumerate(slots):
            top = displayed_value(output, 2 + index * 2, 3 + group_index)
            bottom = displayed_value(output, 3 + index * 2, 3 + group_index)
            middle = (tops[index] + bottoms[index]) / 2
            merged = [area for area in source_merges if area.min_col <= col <= area.max_col
                      and area.min_row <= row + 5 and area.max_row >= row and area.max_col > area.min_col]
            if merged:
                # A merged lecture may be inked only in the central PDF column.
                # Region containment is weaker than group-specific transcription.
                area = max(merged, key=lambda item: item.max_col - item.min_col)
                physical = [bounds[c] for c, _ in groups if c in bounds and area.min_col <= c <= area.max_col]
                if physical:
                    full_left, full_right = min(x[0] for x in physical), max(x[1] for x in physical)
                    near = [w for w in words if full_left < (w[0] + w[2]) / 2 < full_right
                            and tops[index] < (w[1] + w[3]) / 2 < bottoms[index]]
                    regional = tokens(" ".join(w[4] for w in near))
                    if not tokens(top) <= regional or not tokens(bottom) <= regional:
                        findings.append({"type": "merged_region_disagreement", "group": group, "day": day, "time": time})
                    elif top != bottom:
                        middle = (tops[index] + bottoms[index]) / 2
                        pdf_top = tokens(" ".join(w[4] for w in near if (w[1] + w[3]) / 2 < middle))
                        pdf_bottom = tokens(" ".join(w[4] for w in near if (w[1] + w[3]) / 2 >= middle))
                        if tokens(top) <= pdf_top and tokens(bottom) <= pdf_bottom:
                            counts["merged_halves_contained"] += 1
                        else:
                            findings.append({"type": "merged_half_disagreement", "group": group,
                                             "day": day, "time": time})
                counts["merged_region_unverified"] += 1
                continue
            near = [w for w in words if left + 0.4 < (w[0] + w[2]) / 2 < right - 0.4
                    and tops[index] < (w[1] + w[3]) / 2 < bottoms[index]]
            upper = " ".join(w[4] for w in near if (w[1] + w[3]) / 2 < middle)
            lower = " ".join(w[4] for w in near if (w[1] + w[3]) / 2 >= middle)
            divided = any(abs(rule.y0 - middle) < 4.5 and rule.x0 <= left + 1.5
                          and rule.x1 >= right - 1.5 for rule in rules)
            if compare_half(top, bottom, upper, lower, divided):
                counts["exact_local_nonempty" if top or bottom else "exact_local_empty"] += 1
                if divided:
                    counts["exact_divided"] += 1
                continue
            counts["needs_raster_or_span_review"] += 1
            findings.append({"type": "needs_raster_or_span_review", "group": group, "day": day,
                             "time": time, "source": ws.cell(row, col).coordinate,
                             "top": top, "bottom": bottom, "pdf_upper": upper,
                             "pdf_lower": lower, "divided": divided})
    # A repeated class may be typeset once across two slots. Confirm from PDF
    # word union AND absence of a local rule at the intervening time boundary;
    # matching XLSX text alone is not evidence of a span.
    pending = {(item["group"], item["day"], item["time"]): item for item in findings
               if item["type"] == "needs_raster_or_span_review" and item["top"] == item["bottom"] and item["top"]}
    for col, group in groups:
        if col not in bounds:
            continue
        left, right = bounds[col]
        for index, ((row, day, time), (next_row, next_day, next_time)) in enumerate(zip(slots, slots[1:])):
            first = pending.get((group, day, time))
            second = pending.get((group, next_day, next_time))
            if day != next_day or not first or not second or first["top"] != second["top"]:
                continue
            if first["divided"] or second["divided"] or next_row != row + 6:
                continue
            if any(abs(rule.y0 - bottoms[index]) < 1.5 and rule.x0 <= left + 1.5
                   and rule.x1 >= right - 1.5 for rule in rules):
                continue
            pdf_text = " ".join((first["pdf_upper"], first["pdf_lower"],
                                 second["pdf_upper"], second["pdf_lower"]))
            if tokens(first["top"]) != tokens(pdf_text):
                continue
            for item in (first, second):
                item["type"] = "continuous_two_slot_match"
                counts["needs_raster_or_span_review"] -= 1
                counts["continuous_two_slot_match"] += 1

    for item in findings:
        if item["type"] == "needs_raster_or_span_review":
            item["type"] = ("extracted_text_needs_raster" if not item["top"] and not item["bottom"]
                            else "single_week_no_rule_needs_raster" if not item["top"] and item["bottom"]
                            else "unexplained_disagreement")
            counts["needs_raster_or_span_review"] -= 1
            counts[item["type"]] += 1
    counts["positions"] = len(groups) * len(slots)
    return {"counts": dict(counts), "findings": findings,
            "limitations": "Extracted words may be hidden by paint; merged containment is not group proof; raster review required."}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(audit_geometry(args.source, args.pdf, args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
