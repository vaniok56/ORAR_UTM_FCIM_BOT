"""Audit generated schedules against source workbooks and official PDFs."""

from __future__ import annotations

import argparse
import csv
import subprocess
import tempfile
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path

import openpyxl

try:
    from .parser import ScheduleBlock, SourceSegment, _source_horizontal_ranges, parse_workbook, resolve_pairs
except ImportError:  # Direct CLI execution from this directory.
    from parser import ScheduleBlock, SourceSegment, _source_horizontal_ranges, parse_workbook, resolve_pairs


# Covered day-divider residue. The dean repeats these exact cells in every
# revision and they sit between timeslots, so no parser reads them as a class.
# Source owner verified on rendered PDFs 2026-09-29 that they are a bold divider
# line with no visible text. Approved by exact coordinate and text, not by file
# hash, so a new revision keeps the approval but new text does not inherit it.
APPROVED_COVERED_GAPS = {
    (coordinate, "MCE") for coordinate in ("L96", "V96", "W96", "Y96", "Z96", "AA96", "AB96")
} | {("Z182", "sem. ASDN")}


@dataclass(frozen=True)
class AuditFinding:
    status: str
    group: str
    day: str
    time: str
    parity: str
    source: str
    source_text: str
    output_text: str
    reason: str


def normalize(text):
    if text is None:
        return ""
    return "\n".join(
        unicodedata.normalize("NFC", line).strip().replace("0,5", "0.5")
        for line in str(text).splitlines()
        if line.strip()
    )


def lines(text):
    return normalize(text).splitlines()


def compact(text):
    return "".join(normalize(text).casefold().replace("ţ", "ț").split())


def displayed_value(ws, row, column):
    value = ws.cell(row, column).value
    if value is not None:
        return normalize(value)
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return normalize(ws.cell(merged.min_row, merged.min_col).value)
    return ""


def block_values(block: ScheduleBlock):
    values = []
    previous_merge = None
    for value in block.values:
        text = normalize(value.value)
        if not text:
            continue
        if value.merged_range and value.merged_range == previous_merge and values[-1] == text:
            continue
        values.append(text)
        previous_merge = value.merged_range
    return values


def block_lines(block: ScheduleBlock):
    return [line for value in block_values(block) for line in lines(value)]


def segment_text(segment):
    values = []
    previous_merge = None
    for value in segment.values:
        text = normalize(value.value)
        if not text:
            continue
        if value.merged_range and value.merged_range == previous_merge and values[-1] == text:
            continue
        values.append(text)
        previous_merge = value.merged_range
    return "\n".join(values)


def source_neighborhood(group_blocks, index):
    start = max(0, index - 1)
    end = min(len(group_blocks), index + 2)
    values = []
    for block in group_blocks[start:end]:
        if block.day != group_blocks[index].day:
            continue
        values.extend(block_values(block))
    return [line for value in values for line in lines(value)]


def finding(block, status, parity, source_text, output_text, reason):
    return AuditFinding(
        status=status,
        group=block.group,
        day=block.day,
        time=block.time,
        parity={"odd": "ISO-even", "even": "ISO-odd"}.get(parity, parity),
        source=f"{block.values[0].coordinate}:{block.values[-1].coordinate}",
        source_text=source_text,
        output_text=output_text,
        reason=reason,
    )


def audit_schedule(source_path, output_path, official_pdf=None):
    blocks = parse_workbook(source_path)
    pairs = resolve_pairs(blocks)
    groups = list(dict.fromkeys(block.group for block in blocks))
    times = list(dict.fromkeys((block.day, block.time) for block in blocks))
    block_index = {(block.group, block.day, block.time): block for block in blocks}
    pair_index = {(pair.group, pair.day, pair.time): pair for pair in pairs}
    output = openpyxl.load_workbook(output_path, data_only=True).active
    findings = []

    # Parsing only reads six-row slots; report populated separator rows instead
    # of silently treating them as empty schedule content. Covered PDF text can
    # be invisible, so a gap needs explicit source-specific review.
    source = openpyxl.load_workbook(source_path, data_only=True).active
    covered = {block.source_row + offset for block in blocks for offset in range(6)}
    columns = {block.source_column for block in blocks}
    approved = APPROVED_COVERED_GAPS
    for row in range(min(covered), max(covered) + 1):
        if row in covered:
            continue
        for column in sorted(columns):
            cell = source.cell(row, column)
            if cell.value is not None and str(cell.value).strip():
                known = (cell.coordinate, str(cell.value)) in approved
                findings.append(AuditFinding(
                    "approved" if known else "unknown", "", "", "", "source", cell.coordinate,
                    str(cell.value), "", ("covered day-divider text approved by coordinate and text"
                                         if known else "populated row outside six-row timeslots; inspect PDF ink"),
                ))

    if output.max_row != len(times) * 2 + 1 or output.max_column != len(groups) + 2:
        block = blocks[0]
        findings.append(finding(
            block,
            "wrong",
            "both",
            f"{len(times) * 2 + 1} rows x {len(groups) + 2} columns",
            f"{output.max_row} rows x {output.max_column} columns",
            "output dimensions do not match source layout",
        ))

    output_groups = [normalize(output.cell(1, column).value) for column in range(3, output.max_column + 1)]
    if output_groups != groups:
        findings.append(finding(
            blocks[0], "wrong", "both", " | ".join(groups), " | ".join(output_groups), "group headers differ"
        ))

    by_group = {group: [block for block in blocks if block.group == group] for group in groups}
    group_columns = [next(block.source_column for block in blocks if block.group == group) for group in groups]
    for time_index, (day, time) in enumerate(times):
        top_row = 2 + time_index * 2
        if displayed_value(output, top_row, 1) != day:
            findings.append(finding(blocks[0], "wrong", "both", day, displayed_value(output, top_row, 1), "day label differs"))
        if displayed_value(output, top_row, 2) != time:
            findings.append(finding(blocks[0], "wrong", "both", time, displayed_value(output, top_row, 2), "time label differs"))
        for group_index, group in enumerate(groups):
            block = block_index[(group, day, time)]
            pair = pair_index[(group, day, time)]
            odd = displayed_value(output, 2 + time_index * 2, 3 + group_index)
            even = displayed_value(output, 3 + time_index * 2, 3 + group_index)
            source_text = "\n".join(block_lines(block))

            expected_odd = normalize(pair.odd_text)
            expected_even = normalize(pair.even_text)
            if odd != expected_odd:
                findings.append(finding(block, "wrong", "odd", expected_odd, odd, "output differs from resolved upper value"))
            if even != expected_even:
                findings.append(finding(block, "wrong", "even", expected_even, even, "output differs from resolved lower value"))

            # Independent source-border check: an upper/lower divider can be
            # encoded only as the lower cell's top border (no upper bottom).
            if (block.values[3].top_border in {"medium", "thick"}
                    and block.values[2].bottom_border not in {"medium", "thick"}
                    and not (block.values[2].merged_range
                             and block.values[2].merged_range == block.values[3].merged_range)):
                upper = segment_text(SourceSegment(0, 2, block.values[:3]))
                lower = segment_text(SourceSegment(3, 5, block.values[3:]))
                if upper and lower:
                    if odd != upper:
                        findings.append(finding(block, "wrong", "odd", upper, odd, "top-only parity divider disagrees with upper output"))
                    if even != lower:
                        findings.append(finding(block, "wrong", "even", lower, even, "top-only parity divider disagrees with lower output"))

            if "[REVIEW REQUIRED]" in {odd, even}:
                findings.append(finding(block, "wrong", "both", source_text, f"{odd}\n---\n{even}", "placeholder in final output"))

            if pair.status != "auto":
                findings.append(finding(block, "unknown", "both", source_text, f"{odd}\n---\n{even}", "parser could not determine source boundary"))

            joined = any("cross-timeslot" in flag for flag in pair.review_flags)
            source_group_blocks = by_group[group]
            block_position = source_group_blocks.index(block)
            neighborhood = source_neighborhood(source_group_blocks, block_position) if joined else block_lines(block)
            neighborhood_counts = Counter(neighborhood)
            output_union = list(dict.fromkeys(lines(odd) + lines(even)))
            missing = [line for line in block_lines(block) if line not in output_union and line not in lines(
                displayed_value(output, max(2, 2 + (time_index - 1) * 2), 3 + group_index)
                + "\n"
                + displayed_value(output, min(output.max_row, 3 + (time_index + 1) * 2), 3 + group_index)
            )]
            if missing:
                findings.append(finding(block, "missing", "both", source_text, f"{odd}\n---\n{even}", f"source lines absent: {missing}"))

            extra = [line for line in output_union if line not in neighborhood]
            if extra:
                findings.append(finding(block, "extra", "both", source_text, f"{odd}\n---\n{even}", f"output lines not in source neighborhood: {extra}"))

            for parity, text in (("odd", odd), ("even", even)):
                duplicates = [
                    line for line, count in Counter(lines(text)).items()
                    if count > max(1, neighborhood_counts[line])
                ]
                if duplicates:
                    findings.append(finding(block, "extra", parity, source_text, text, f"duplicated output lines: {duplicates}"))

            segments = [segment for segment in block.segments if segment_text(segment)]
            shape = tuple((segment.start_offset, segment.end_offset) for segment in segments)
            if shape == ((0, 2), (3, 5)) and not joined:
                expected_odd = segment_text(segments[0])
                expected_even = segment_text(segments[1])
                if odd != expected_odd:
                    findings.append(finding(block, "wrong", "odd", expected_odd, odd, "upper output disagrees with explicit source border"))
                if even != expected_even:
                    findings.append(finding(block, "wrong", "even", expected_even, even, "lower output disagrees with explicit source border"))
            elif shape == ((0, 2),) and not joined and odd != segment_text(segments[0]):
                findings.append(finding(block, "wrong", "odd", segment_text(segments[0]), odd, "upper output disagrees with upper source segment"))
            elif shape == ((3, 5),) and not joined and even != segment_text(segments[0]):
                findings.append(finding(block, "wrong", "even", segment_text(segments[0]), even, "lower output disagrees with lower source segment"))

            for parity, offsets, output_row, output_text in (
                ("odd", range(3), top_row, odd),
                ("even", range(3, 6), top_row + 1, even),
            ):
                if not output_text:
                    continue
                source_ranges = _source_horizontal_ranges(block, group_columns)[0 if parity == "odd" else 1]
                for first, last in source_ranges:
                    first_column, last_column = first + 3, last + 3
                    if not any(
                        merged.min_row <= output_row <= merged.max_row
                        and merged.min_col <= first_column
                        and merged.max_col >= last_column
                        for merged in output.merged_cells.ranges
                    ):
                        findings.append(finding(
                            block,
                            "wrong",
                            parity,
                            f"horizontal merge {first_column}:{last_column}",
                            "unmerged",
                            "source-backed horizontal merge missing",
                        ))

    if official_pdf:
        with tempfile.NamedTemporaryFile(suffix=".txt") as extracted:
            try:
                subprocess.run(
                    ["pdftotext", "-enc", "UTF-8", "-raw", "-nopgbrk", str(official_pdf), extracted.name],
                    check=True, capture_output=True, timeout=30,
                )
            except subprocess.TimeoutExpired as error:
                raise ValueError("PDF text extraction exceeded 30 seconds") from error
            pdf_text = compact(Path(extracted.name).read_text(encoding="utf-8", errors="replace"))
        unique_values = list(dict.fromkeys(value for block in blocks for value in block_values(block)))
        for value in unique_values:
            variants = [compact(value)]
            if compact(value).startswith("aula"):
                short = normalize(value)[5:]
                variants.extend([compact(short), compact(short.replace("Henri Coandă", ""))])
            if not any(variant and variant in pdf_text for variant in variants):
                block = next(block for block in blocks if value in block_values(block))
                findings.append(finding(block, "unknown", "source", value, "", "source value not found by PDF text extraction"))

    total_cells = len(groups) * len(times) * 2
    return findings, {
        "groups": len(groups),
        "timeslots": len(times),
        "source_blocks": len(blocks),
        "output_parity_cells": total_cells,
    }


def write_csv(findings, path):
    with Path(path).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=AuditFinding.__dataclass_fields__)
        writer.writeheader()
        writer.writerows({key: csv_cell(value) for key, value in asdict(item).items()} for item in findings)


def csv_cell(value):
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")) else value


def write_report(findings, stats, path):
    counts = Counter(item.status for item in findings)
    failed_positions = set()
    for item in findings:
        if item.parity in {"ISO-even", "ISO-odd"}:
            failed_positions.add((item.group, item.day, item.time, item.parity))
        elif item.parity == "both":
            failed_positions.add((item.group, item.day, item.time, "ISO-even"))
            failed_positions.add((item.group, item.day, item.time, "ISO-odd"))
    passed = max(0, stats["output_parity_cells"] - len(failed_positions))
    report = [
        "# Schedule audit",
        "",
        f"- Groups: {stats['groups']}",
        f"- Timeslots: {stats['timeslots']}",
        f"- Source blocks: {stats['source_blocks']}",
        f"- Output parity cells: {stats['output_parity_cells']}",
        f"- Output parity cells without findings: {passed} (source-level findings excluded)",
        f"- Wrong: {counts['wrong']}",
        f"- Missing: {counts['missing']}",
        f"- Extra: {counts['extra']}",
        f"- Unknown: {counts['unknown']}",
        f"- Approved covered source values: {counts['approved']}",
        "",
        "Scope: source/output and PDF text presence only. Does not verify rendered PDF geometry or covered text.",
        "",
        "Safe: " + ("yes" if all(item.status == "approved" for item in findings) else "no"),
        "",
    ]
    Path(path).write_text("\n".join(report), encoding="utf-8")


def main():
    cli = argparse.ArgumentParser()
    cli.add_argument("--source", required=True)
    cli.add_argument("--output", required=True)
    cli.add_argument("--official-pdf")
    cli.add_argument("--csv", required=True)
    cli.add_argument("--report", required=True)
    args = cli.parse_args()
    findings, stats = audit_schedule(args.source, args.output, args.official_pdf)
    write_csv(findings, args.csv)
    write_report(findings, stats, args.report)
    counts = Counter(item.status for item in findings)
    print(f"Wrong: {counts['wrong']} Missing: {counts['missing']} Extra: {counts['extra']} Unknown: {counts['unknown']} Approved: {counts['approved']}")
    raise SystemExit(1 if any(item.status != "approved" for item in findings) else 0)


if __name__ == "__main__":
    main()
