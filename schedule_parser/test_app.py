import json
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from subprocess import CompletedProcess
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from app import (
    Cancelled,
    ScheduleInput,
    ScheduleMetadata,
    ScheduleResult,
    add_schedule_files,
    ask_file,
    ask_version,
    copy_sources,
    create_run_directory,
    inspect_pdf,
    inspect_xlsx,
    pair_mismatches,
    parse_path,
    parse_paths,
    parse_pdf_version,
    parse_schedule_title,
    pdf_tools_available,
    process_session,
    read_pdf_text,
    render_pdf,
    revision_key,
    run_new_session,
    run_wizard,
    schedule_label,
    sha256,
    validate_existing,
    write_session_files,
)
from parser import parse_workbook, write_candidate_review_workbook


DOWNLOADS = Path.home() / "Downloads"
CURRENT = {
    "II": (
        DOWNLOADS / "Anul_II_Semestrul_III-2.xlsx",
        DOWNLOADS / "anul_ii_semestrul_iii-5-2.pdf",
    ),
    "III": (
        DOWNLOADS / "Anul_III_Semestrul_V.xlsx",
        DOWNLOADS / "anul_iii_semestrul_v-4-2.pdf",
    ),
    "IV": (
        DOWNLOADS / "Anul_IV_Semestrul_VII.xlsx",
        DOWNLOADS / "anul_iv_semestrul_vii-6-2.pdf",
    ),
}
CURRENT_AVAILABLE = all(path.is_file() for pair in CURRENT.values() for path in pair)


class AppTests(unittest.TestCase):
    def test_pdf_text_extraction_requests_utf8(self):
        with patch("app.subprocess.run", return_value=CompletedProcess([], 0, stdout="Orar")) as run:
            self.assertEqual(read_pdf_text(Path("official.pdf")), "Orar")
        self.assertEqual(
            run.call_args.args[0],
            ["pdftotext", "-enc", "UTF-8", "-raw", "-nopgbrk", "official.pdf", "-"],
        )

    def test_pdf_preview_uses_poppler_cairo_renderer(self):
        with patch("app.subprocess.run") as run:
            render_pdf(Path("official.pdf"), Path("official.png"))
        self.assertEqual(run.call_args.args[0][0], "pdftocairo")
        self.assertIn("-singlefile", run.call_args.args[0])
        self.assertIn("-png", run.call_args.args[0])

    def test_parse_path_accepts_quoted_dragged_path(self):
        self.assertEqual(parse_path("'/tmp/My Schedule.xlsx'"), Path("/tmp/My Schedule.xlsx").resolve())

    def test_parse_path_rejects_multiple_paths(self):
        with self.assertRaises(ValueError):
            parse_path("/tmp/one.xlsx /tmp/two.xlsx")

    def test_parse_paths_accepts_multiple_escaped_paths(self):
        self.assertEqual(
            parse_paths(r"/tmp/Year\ I.xlsx '/tmp/Year I.pdf'"),
            [Path("/tmp/Year I.xlsx").resolve(), Path("/tmp/Year I.pdf").resolve()],
        )

    def test_optional_file_accepts_empty_input(self):
        self.assertIsNone(ask_file("PDF", ".pdf", optional=True, input_fn=lambda _: ""))

    def test_required_file_rejects_empty_input(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "schedule.xlsx"
            path.touch()
            answers = iter(("", str(path)))
            messages = []
            self.assertEqual(
                ask_file("XLSX", ".xlsx", input_fn=lambda _: next(answers), output_fn=messages.append),
                path.resolve(),
            )
            self.assertIn("Path is required.", messages)

    def test_file_prompt_rejects_wrong_suffix_then_accepts(self):
        with tempfile.TemporaryDirectory() as directory:
            wrong = Path(directory) / "schedule.pdf"
            right = Path(directory) / "schedule.xlsx"
            wrong.touch()
            right.touch()
            answers = iter((str(wrong), str(right)))
            messages = []
            self.assertEqual(
                ask_file("XLSX", ".xlsx", input_fn=lambda _: next(answers), output_fn=messages.append),
                right.resolve(),
            )
            self.assertIn("Expected a .xlsx file.", messages)

    def test_file_prompt_identifies_folder_path(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            file = root / "schedule.xlsx"
            file.touch()
            answers = iter((str(root), str(file)))
            messages = []
            self.assertEqual(
                ask_file("XLSX", ".xlsx", input_fn=lambda _: next(answers), output_fn=messages.append),
                file.resolve(),
            )
            self.assertIn("Expected a .xlsx file, but this path is a folder.", messages)

    def test_file_prompt_can_cancel(self):
        with self.assertRaises(Cancelled):
            ask_file("XLSX", ".xlsx", input_fn=lambda _: "c")

    def test_schedule_title_parser(self):
        metadata = parse_schedule_title(
            "ACTIVITĂŢILOR DIDACTICE ÎN ANUL UNIVERSITAR 2026/2027, ANUL II, SEMESTRUL III"
        )
        self.assertEqual((metadata.academic_year, metadata.study_year, metadata.semester), ("2026/2027", "II", "III"))

    def test_pdf_version_parser_ignores_download_duplicate_suffix(self):
        cases = {
            "anul_iv_semestrul_vii-6.pdf": 6,
            "anul_iii_semestrul_v-4-2.pdf": 4,
            "anul_ii_semestrul_iii-12-7.pdf": 12,
            "anul_ii_semestrul_iii.pdf": None,
            "anul_ii_semestrul_iii-100.pdf": None,
            "anul_ii_semestrul_iii-00.pdf": None,
        }
        for name, expected in cases.items():
            with self.subTest(name=name):
                self.assertEqual(parse_pdf_version(Path(name)), expected)

    def test_detected_pdf_version_can_be_used_or_overridden(self):
        pdf = Path("anul_iv_semestrul_vii-6-2.pdf")
        self.assertEqual(ask_version(pdf, input_fn=lambda _: "", output_fn=lambda _: None), 6)
        self.assertEqual(ask_version(pdf, input_fn=lambda _: "12", output_fn=lambda _: None), 12)

    def test_missing_pdf_requires_explicit_version_or_final(self):
        answers = iter(("", "0", "100", "09"))
        messages = []
        self.assertEqual(ask_version(None, lambda _: next(answers), messages.append), 9)
        self.assertEqual(len(messages), 3)
        self.assertEqual(ask_version(None, lambda _: "final", lambda _: None), 0)

    def test_pair_mismatch_reports_exact_fields(self):
        xlsx = ScheduleMetadata("2026/2027", "II", "III")
        pdf = ScheduleMetadata("2026/2027", "III", "V")
        self.assertEqual(pair_mismatches(xlsx, pdf), ["study_year", "semester"])

    def test_revision_key_ignores_file_names(self):
        metadata = ScheduleMetadata("2026/2027", "II", "III")
        first = ScheduleInput(Path("first.xlsx"), None, None, metadata, 1)
        second = ScheduleInput(Path("second.xlsx"), None, None, metadata, 2)
        self.assertEqual(revision_key(first), revision_key(second))

    def test_schedule_label_distinguishes_academic_years(self):
        first = schedule_label(ScheduleMetadata("2025/2026", "II", "III"))
        second = schedule_label(ScheduleMetadata("2026/2027", "II", "III"))
        self.assertNotEqual(first, second)

    def test_run_directory_is_unique_and_immutable(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            now = datetime(2026, 9, 1, 12, 0, 0, 123456)
            created = create_run_directory(root, now)
            self.assertTrue(created.is_dir())
            with self.assertRaises(FileExistsError):
                create_run_directory(root, now)

    def test_sources_are_copied_and_hashes_match(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            source.write_bytes(b"xlsx source")
            output = root / "year_II"
            output.mkdir()
            item = ScheduleInput(source, None, None, ScheduleMetadata("2026/2027", "II", "III"), 7)
            copied, manifest = copy_sources(item, output)
            self.assertEqual(copied.xlsx.read_bytes(), source.read_bytes())
            self.assertEqual(manifest["xlsx_sha256"], sha256(source))
            self.assertEqual(manifest["validation_scope"], "xlsx-only")
            self.assertEqual(manifest["version"], 7)

    @unittest.skipUnless(CURRENT_AVAILABLE, "current schedule files are unavailable")
    def test_candidate_review_works_without_transform_reference(self):
        blocks = parse_workbook(CURRENT["II"][0])
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "candidate.xlsx"
            write_candidate_review_workbook(blocks, None, output)
            ws = openpyxl.load_workbook(output, data_only=True).active
            self.assertGreater(ws.max_row, 1)
            self.assertTrue(all(ws.cell(row, 14).value is None for row in range(2, ws.max_row + 1)))
            self.assertTrue(all(ws.cell(row, 15).value is None for row in range(2, ws.max_row + 1)))

    def test_batch_rejects_pdf_without_matching_xlsx_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            xlsx = root / "source.xlsx"
            pdf = root / "official.pdf"
            xlsx.touch()
            pdf.touch()
            with (
                patch("app.inspect_xlsx", return_value=ScheduleMetadata("2026/2027", "II", "III", ("TI-251",))),
                patch("app.inspect_pdf", return_value=ScheduleMetadata("2026/2027", "III", "V")),
                patch("app.pdf_tools_available", return_value=True),
            ):
                with self.assertRaisesRegex(ValueError, "PDF has no matching XLSX"):
                    add_schedule_files(lambda _: f"{xlsx} {pdf}", lambda _: None)

    def test_batch_pairs_multiple_files_in_any_order(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths = [
                root / "year_ii-18.pdf",
                root / "year_i-27.pdf",
                root / "year_i.xlsx",
                root / "year_ii.xlsx",
            ]
            for path in paths:
                path.touch()
            metadata = {
                "year_i.xlsx": ScheduleMetadata("2026/2027", "I", "I", ("SI-261",)),
                "year_ii.xlsx": ScheduleMetadata("2026/2027", "II", "III", ("TI-251",)),
                "year_i-27.pdf": ScheduleMetadata("2026/2027", "I", "I"),
                "year_ii-18.pdf": ScheduleMetadata("2026/2027", "II", "III"),
            }
            answers = iter((" ".join(str(path) for path in paths), "27", "18"))
            with (
                patch("app.inspect_xlsx", side_effect=lambda path: metadata[path.name]),
                patch("app.inspect_pdf", side_effect=lambda path: metadata[path.name]),
                patch("app.pdf_tools_available", return_value=True),
            ):
                items = add_schedule_files(lambda _: next(answers), lambda _: None)
            self.assertEqual([item.metadata.study_year for item in items], ["I", "II"])
            self.assertEqual([item.pdf.name for item in items], ["year_i-27.pdf", "year_ii-18.pdf"])
            self.assertEqual([item.version for item in items], [27, 18])

    def test_batch_accepts_xlsx_only_with_explicit_final_version(self):
        with tempfile.TemporaryDirectory() as directory:
            xlsx = Path(directory) / "year_i.xlsx"
            xlsx.touch()
            answers = iter((str(xlsx), "final"))
            with patch(
                "app.inspect_xlsx",
                return_value=ScheduleMetadata("2026/2027", "I", "I", ("SI-261",)),
            ):
                items = add_schedule_files(lambda _: next(answers), lambda _: None)
            self.assertEqual(len(items), 1)
            self.assertIsNone(items[0].pdf)
            self.assertEqual(items[0].version, 0)

    def test_wizard_uses_only_entered_paths(self):
        metadata = ScheduleMetadata("2026/2027", "II", "III", ("TI-251",))
        item = ScheduleInput(Path("source.xlsx"), None, None, metadata, 1)
        result = ScheduleResult("year_II_semester_III", "xlsx-only", 1, "2026-09-01", 1, 70, 0, 0, 0, 0, 0, "SAFE, XLSX ONLY", Path("/tmp/run/year"))
        answers = iter(("1", "1", "3", "", "", "q"))
        with (
            patch("app.add_schedule_files", return_value=[item]),
            patch("app.process_session", return_value=(Path("/tmp/run"), [result], [])),
            patch("pathlib.Path.glob") as glob,
            patch("os.walk") as walk,
        ):
            exit_code = run_wizard(lambda _: next(answers), lambda _: None, open_fn=lambda _: None)
        self.assertEqual(exit_code, 0)
        glob.assert_not_called()
        walk.assert_not_called()

    def test_wizard_keeps_failed_validation_status_after_returning_to_menu(self):
        answers = iter(("2", "q"))
        with patch("app.validate_existing", return_value=1):
            self.assertEqual(run_wizard(lambda _: next(answers), lambda _: None), 1)

    def test_cancel_exits_without_creating_run(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            answers = iter(("1", "c", "q"))
            self.assertEqual(run_wizard(lambda _: next(answers), lambda _: None, root), 2)
            self.assertEqual(list(root.iterdir()), [])

    def test_duplicate_revision_can_replace_selected_item(self):
        metadata = ScheduleMetadata("2026/2027", "II", "III")
        first = ScheduleInput(Path("first.xlsx"), None, None, metadata, 1)
        second = ScheduleInput(Path("second.xlsx"), None, None, metadata, 2)
        result = ScheduleResult("year_II_semester_III", "xlsx-only", 2, "2026-09-01", 1, 70, 0, 0, 0, 0, 0, "SAFE, XLSX ONLY", Path("/tmp/run/year"))
        answers = iter(("1", "1", "1", "3", "", "q"))
        with (
            patch("app.add_schedule_files", side_effect=([first], [second])),
            patch("app.process_session", return_value=(Path("/tmp/run"), [result], [])) as process,
        ):
            self.assertEqual(run_new_session(lambda _: next(answers), lambda _: None), (0, True))
        self.assertEqual(process.call_args.args[0], [second])

    def test_failed_audit_returns_nonzero(self):
        metadata = ScheduleMetadata("2026/2027", "II", "III")
        item = ScheduleInput(Path("source.xlsx"), None, None, metadata, 1)
        result = ScheduleResult("label", "xlsx-only", 1, "2026-09-01", 1, 70, 0, 1, 0, 0, 0, "FAILED", Path("/tmp/run/year"))
        answers = iter(("1", "3", "", "q"))
        with (
            patch("app.add_schedule_files", return_value=[item]),
            patch("app.process_session", return_value=(Path("/tmp/run"), [result], [])),
        ):
            self.assertEqual(run_new_session(lambda _: next(answers), lambda _: None), (1, True))

    def test_open_failure_keeps_final_action_menu_active(self):
        metadata = ScheduleMetadata("2026/2027", "II", "III")
        item = ScheduleInput(Path("source.xlsx"), None, None, metadata, 1)
        result = ScheduleResult("label", "xlsx-only", 1, "2026-09-01", 1, 70, 0, 0, 0, 0, 0, "SAFE, XLSX ONLY", Path("/tmp/run/year"))
        answers = iter(("1", "3", "", "1", "q"))
        messages = []

        def fail_open(_):
            raise OSError("boom")

        with (
            patch("app.add_schedule_files", return_value=[item]),
            patch("app.process_session", return_value=(Path("/tmp/run"), [result], [])),
        ):
            outcome = run_new_session(
                lambda _: next(answers), messages.append, open_fn=fail_open
            )
        self.assertEqual(outcome, (0, True))
        self.assertTrue(any("Could not open file" in message for message in messages))

    def test_failed_schedule_keeps_verified_source_manifest(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            source.write_bytes(b"source")
            item = ScheduleInput(source, None, None, ScheduleMetadata("2026/2027", "II", "III"), 1)
            with patch("app.parse_workbook", side_effect=ValueError("broken workbook")):
                run, results, failures = process_session([item], root / "runs")
            manifest = json.loads((run / "session.json").read_text(encoding="utf-8"))
            self.assertEqual(results, [])
            self.assertEqual(len(failures), 1)
            self.assertEqual(len(manifest["schedules"]), 1)
            self.assertTrue((run / manifest["schedules"][0]["copied_xlsx"]).is_file())

    def test_validate_existing_rejects_mismatched_pdf(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            generated = root / "generated.xlsx"
            pdf = root / "official.pdf"
            for path in (source, generated, pdf):
                path.touch()
            answers = iter((str(source), str(generated), str(pdf)))
            with (
                patch("app.pdf_tools_available", return_value=True),
                patch("app.inspect_xlsx", return_value=ScheduleMetadata("2026/2027", "II", "III")),
                patch("app.inspect_pdf", return_value=ScheduleMetadata("2026/2027", "III", "V")),
            ):
                result = validate_existing(lambda _: next(answers), lambda _: None)
            self.assertEqual(result, 1)
            self.assertFalse((root / "audit.csv").exists())

    def test_xlsx_only_summary_does_not_claim_pdf_validation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            result = ScheduleResult(
                "year_II_semester_III", "xlsx-only", 5, "2026-09-01", 26, 1820, 47,
                0, 0, 0, 0, "SAFE, XLSX ONLY", root / "year_II",
            )
            write_session_files(root, [result], [], [])
            summary = (root / "SUMMARY.md").read_text(encoding="utf-8")
            self.assertIn("SAFE, XLSX ONLY", summary)
            self.assertNotIn("SAFE + PDF", summary)
            self.assertIn("Version: 5", summary)
            self.assertIn("Generated date: 2026-09-01", summary)

    @unittest.skipUnless(CURRENT_AVAILABLE and pdf_tools_available(), "current schedule files or PDF tools unavailable")
    def test_older_three_year_session_passes_with_coordinate_approved_gaps(self):
        items = []
        for xlsx, pdf in CURRENT.values():
            xlsx_metadata = inspect_xlsx(xlsx)
            self.assertEqual(pair_mismatches(xlsx_metadata, inspect_pdf(pdf)), [])
            items.append(ScheduleInput(xlsx, pdf, None, xlsx_metadata, parse_pdf_version(pdf)))
        with tempfile.TemporaryDirectory() as directory:
            progress = []
            started_at = datetime(2026, 9, 1, 14, 30)
            run, results, failures = process_session(items, Path(directory), progress.append, started_at)
            self.assertEqual(failures, [])
            self.assertEqual(len(results), 3)
            for item, result in zip(items, results):
                # Covered day-divider cells are approved by coordinate and text,
                # so no year blocks on them any more.
                self.assertEqual((result.wrong, result.missing, result.extra, result.unknown), (0, 0, 0, 0))
                self.assertEqual(result.verdict, "SOURCE + PDF TEXT PASS (GEOMETRY NOT CHECKED)")
                audit_rows = [row for row in (result.output / "audit.csv").read_text(encoding="utf-8").splitlines() if row]
                if item.metadata.study_year in {"II", "III"}:
                    self.assertTrue(all(row.startswith("approved,") for row in audit_rows[1:]))
                else:
                    self.assertEqual(audit_rows[1:], [])
                self.assertEqual(result.version, item.version)
                self.assertEqual(result.generated_date, "2026-09-01")
                self.assertTrue((result.output / "sources" / "dean.xlsx").is_file())
                self.assertTrue((result.output / "sources" / "official.pdf").is_file())
                audit_rows = [row for row in (result.output / "audit.csv").read_text(encoding="utf-8").splitlines() if row]
                self.assertTrue(all(row.startswith("approved,") for row in audit_rows[1:]))
                ws = openpyxl.load_workbook(result.output / "final_schedule.xlsx", data_only=True).active
                self.assertEqual(ws["A1"].value, item.version)
                self.assertEqual(ws["B1"].value.date(), started_at.date())
                self.assertEqual(ws["B1"].number_format, "dd/mm/yyyy")
            self.assertTrue((run / "SUMMARY.md").is_file())
            self.assertTrue((run / "summary.csv").is_file())
            self.assertTrue((run / "session.json").is_file())
            manifest = json.loads((run / "session.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["created_at"], started_at.isoformat())
            self.assertEqual([entry["version"] for entry in manifest["schedules"]], [5, 4, 6])
            self.assertEqual({entry["generated_date"] for entry in manifest["schedules"]}, {"2026-09-01"})
            self.assertTrue(any("Copied and verified sources" in message for message in progress))
            self.assertEqual(sum("Source/PDF text checks passed" in message for message in progress), 3)

    @unittest.skipUnless(CURRENT_AVAILABLE, "current schedule files are unavailable")
    def test_older_xlsx_only_session_passes_with_coordinate_approved_gaps(self):
        xlsx = CURRENT["II"][0]
        item = ScheduleInput(xlsx, None, None, inspect_xlsx(xlsx), 8)
        with tempfile.TemporaryDirectory() as directory:
            run, results, failures = process_session([item], Path(directory))
            self.assertEqual(failures, [])
            self.assertEqual(results[0].verdict, "SOURCE PASS (NO PDF)")
            self.assertEqual(results[0].unknown, 0)
            self.assertIn("approved", (results[0].output / "audit.csv").read_text(encoding="utf-8"))
            self.assertFalse((results[0].output / "sources" / "official.pdf").exists())
            self.assertFalse((results[0].output / "official.png").exists())
            self.assertIn("xlsx-only", (run / "SUMMARY.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
